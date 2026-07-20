#!/usr/bin/env python3
"""Evaluate human labels from the 200-row real Prom review workbook."""

from __future__ import annotations

import argparse
from collections import Counter
import hashlib
import json
from pathlib import Path
import sys
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend" / "src"))

from metis.pricing import (  # noqa: E402
    ProductTier,
    TierValidationCase,
    classify_tier,
    evaluate_tier_validation,
    load_approved_brand_rules,
)


DEFAULT_WORKBOOK = (
    ROOT
    / "outputs"
    / "metis_action_plan_20260719"
    / "METIS_200_REAL_LISTING_REVIEW.xlsx"
)
DEFAULT_BRANDS = ROOT / "backend" / "config" / "brands.yaml"
REQUIRED_LABEL_FIELDS = (
    "same_part_gold",
    "condition_gold",
    "tier_gold",
    "include_in_target_gold",
    "reviewer",
)


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _read_cases(
    workbook_path: Path,
    *,
    sheet_name: str,
    brand_rules,
) -> tuple[list[TierValidationCase], dict[str, int], list[str]]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=False)
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"Workbook has no sheet named {sheet_name!r}")
    sheet = workbook[sheet_name]
    headers = {
        _text(cell.value): index
        for index, cell in enumerate(sheet[4], start=1)
        if _text(cell.value)
    }
    required_source_fields = {
        "source_listing_id",
        "url",
        "brand",
        "title",
        "description",
        *REQUIRED_LABEL_FIELDS,
    }
    missing_headers = sorted(required_source_fields - set(headers))
    if missing_headers:
        raise ValueError(f"Workbook is missing headers: {missing_headers}")

    cases: list[TierValidationCase] = []
    counts = {
        "source_rows": 0,
        "unique_source_rows": 0,
        "missing_source_slots": 0,
        "pending_label_rows": 0,
        "labeled_rows": 0,
        "included_metric_rows": 0,
        "excluded_metric_rows": 0,
    }
    errors: list[str] = []
    source_ids: list[str] = []
    source_urls: list[str] = []
    for row_number, row in enumerate(
        sheet.iter_rows(min_row=5, values_only=True), start=5
    ):
        values = {
            header: row[column - 1] if column <= len(row) else None
            for header, column in headers.items()
        }
        source_listing_id = _text(values.get("source_listing_id"))
        if not source_listing_id:
            counts["missing_source_slots"] += 1
            continue
        counts["source_rows"] += 1
        source_ids.append(source_listing_id)
        source_urls.append(_text(values.get("url")))
        labels = {field: _text(values.get(field)) for field in REQUIRED_LABEL_FIELDS}
        if not all(labels.values()):
            counts["pending_label_rows"] += 1
            continue
        counts["labeled_rows"] += 1

        same_part = labels["same_part_gold"].casefold()
        condition = labels["condition_gold"].casefold()
        include = labels["include_in_target_gold"].casefold()
        if same_part not in {"yes", "no", "uncertain"}:
            errors.append(
                f"ROW_{row_number}_INVALID_SAME_PART:{labels['same_part_gold']!r}"
            )
            continue
        if condition not in {"new", "used", "refurbished", "unknown", "conflict"}:
            errors.append(
                f"ROW_{row_number}_INVALID_CONDITION:{labels['condition_gold']!r}"
            )
            continue
        if include not in {"yes", "no", "uncertain"}:
            errors.append(
                f"ROW_{row_number}_INVALID_INCLUDE:{labels['include_in_target_gold']!r}"
            )
            continue
        try:
            gold_tier = ProductTier(labels["tier_gold"])
        except ValueError:
            errors.append(
                f"ROW_{row_number}_INVALID_GOLD_TIER:{labels['tier_gold']!r}"
            )
            continue
        if include != "yes":
            counts["excluded_metric_rows"] += 1
            continue
        if same_part != "yes":
            errors.append(f"ROW_{row_number}_INCLUDE_REQUIRES_SAME_PART_YES")
            continue
        if condition in {"used", "refurbished", "conflict"}:
            errors.append(f"ROW_{row_number}_INCLUDE_REJECTS_{condition.upper()}")
            continue
        prediction = classify_tier(
            brand=_text(values.get("brand")) or None,
            title=_text(values.get("title")),
            description=_text(values.get("description")) or None,
            brand_tiers=brand_rules.tiers,
        )
        cases.append(
            TierValidationCase(
                source_listing_id=source_listing_id,
                gold_tier=gold_tier,
                predicted_tier=prediction.tier,
            )
        )
        counts["included_metric_rows"] += 1
    counts["unique_source_rows"] = len(set(source_ids))
    duplicate_ids = sorted(
        source_id
        for source_id, count in Counter(source_ids).items()
        if count > 1
    )
    duplicate_urls = sorted(
        source_url
        for source_url, count in Counter(source_urls).items()
        if source_url and count > 1
    )
    errors.extend(f"DUPLICATE_SOURCE_LISTING_ID:{value}" for value in duplicate_ids)
    errors.extend(f"DUPLICATE_SOURCE_URL:{value}" for value in duplicate_urls)
    workbook.close()
    return cases, counts, errors


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--brands", type=Path, default=DEFAULT_BRANDS)
    parser.add_argument("--sheet", default="Разметка 200")
    parser.add_argument("--output", type=Path)
    parser.add_argument("--min-labeled", type=int, default=200)
    parser.add_argument("--min-included", type=int, default=1)
    parser.add_argument("--representative", action="store_true")
    parser.add_argument("--approved-by")
    parser.add_argument("--allow-blocked", action="store_true")
    args = parser.parse_args()

    workbook_raw = args.workbook.read_bytes()
    rules = load_approved_brand_rules(args.brands)
    cases, workbook_counts, workbook_errors = _read_cases(
        args.workbook,
        sheet_name=args.sheet,
        brand_rules=rules,
    )
    report = evaluate_tier_validation(
        cases,
        min_labeled=args.min_included,
        representative=args.representative,
        approved_by=args.approved_by,
        domain_policy_approved=rules.domain_policy_approved,
    )
    report["input"] = {
        "workbook": str(args.workbook.resolve()),
        "workbook_sha256": hashlib.sha256(workbook_raw).hexdigest(),
        "sheet": args.sheet,
        "brand_dictionary": rules.source_path,
        "brand_dictionary_sha256": rules.source_sha256,
        "brand_dataset_id": rules.dataset_id,
    }
    report["workbook_counts"] = workbook_counts
    report_counts = report["counts"]
    if isinstance(report_counts, dict):
        report_counts["reviewed"] = workbook_counts["labeled_rows"]
        report_counts["excluded_from_metrics"] = workbook_counts[
            "excluded_metric_rows"
        ]
        report_counts["minimum_reviewed_required"] = args.min_labeled
        report_counts["minimum_included_required"] = args.min_included
    if workbook_counts["source_rows"] < args.min_labeled:
        report["activation_blockers"] = [
            *report["activation_blockers"],
            "SOURCE_ROWS_BELOW_REVIEW_MINIMUM:"
            f"{workbook_counts['source_rows']}/{args.min_labeled}",
        ]
    if workbook_counts["labeled_rows"] < args.min_labeled:
        report["activation_blockers"] = [
            *report["activation_blockers"],
            "REVIEWED_ROWS_BELOW_MINIMUM:"
            f"{workbook_counts['labeled_rows']}/{args.min_labeled}",
        ]
    if workbook_errors:
        report["contract_errors"] = [
            *report["contract_errors"],
            *workbook_errors,
        ]
        report["validation_gate"] = "FAIL"
    elif report["activation_blockers"]:
        report["validation_gate"] = "BLOCKED"

    rendered = json.dumps(report, indent=2, sort_keys=True) + "\n"
    if args.output is not None:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(rendered, encoding="utf-8")
    print(rendered, end="")
    if report["validation_gate"] == "PASS" or args.allow_blocked:
        return 0
    return 1


if __name__ == "__main__":
    raise SystemExit(main())
