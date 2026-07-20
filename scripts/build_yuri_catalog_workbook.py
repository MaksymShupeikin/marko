#!/usr/bin/env python3
"""Build a safe, operator-friendly Marko workbook from Yuri's Prom.ua export."""

from __future__ import annotations

import argparse
from collections import Counter, defaultdict
import hashlib
import json
from math import sqrt
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from marko.services.cost_privacy import is_raw_cost_label
from marko.services.xlsx_catalog import normalize_identifier, parse_catalog_xlsx


SOURCE_SHEET = "Export Products Sheet"
GENERATED_SHEETS = ("Инструкция", "Ввод Юрия", "Проверить OE")
INPUT_HEADERS = (
    "SKU",
    "OE",
    "Название",
    "Категория",
    "Бренд",
    "Цена",
    "Валюта",
    "Остаток",
    "Наличие",
    "Ссылка на товар",
    "MPN",
    "Описание",
    "Поисковые запросы",
    "Статус товара",
    "Дней на складе",
    "Себестоимость",
    "Продажи за 30 дней",
    "Продажи за 60 дней",
    "Продажи за 90 дней",
    "Дней с последней продажи",
    "Средние продажи в месяц",
    "План продаж в месяц",
    "Просмотры за 30 дней",
    "Конверсия",
    "Приоритет",
    "Комментарий Юрия",
    "Исходная строка",
    "Маркер наличия Prom.ua",
)

NAVY = "17324D"
BLUE = "0066CC"
WHITE = "FFFFFF"
TEXT = "23313F"
MUTED = "617182"
PALE_BLUE = "EAF2F8"
PALE_YELLOW = "FFF4CC"
PALE_ORANGE = "FCE4D6"
PALE_RED = "FDE9E7"
PALE_GREEN = "E2F0D9"
PALE_GRAY = "EEF1F4"
BORDER = Side(style="thin", color="CBD5DF")


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, float):
        return format(value, ".15g")
    return str(value).strip()


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _wilson(successes: int, total: int, z: float = 1.959963984540054) -> dict[str, float]:
    if total == 0:
        return {"point": 0.0, "lower": 0.0, "upper": 0.0}
    p = successes / total
    z2 = z * z
    denominator = 1 + z2 / total
    center = (p + z2 / (2 * total)) / denominator
    margin = z * sqrt(p * (1 - p) / total + z2 / (4 * total * total)) / denominator
    return {"point": p, "lower": center - margin, "upper": center + margin}


def _source_rows(sheet) -> tuple[list[str], list[tuple[Any, ...]]]:
    iterator = sheet.iter_rows(values_only=True)
    headers = [_text(value) for value in next(iterator)]
    rows = [tuple(row) for row in iterator if any(value not in (None, "") for value in row)]
    return headers, rows


def _sheet_digest(sheet) -> str:
    digest = hashlib.sha256()
    for row in sheet.iter_rows(values_only=True):
        payload = json.dumps(row, ensure_ascii=False, default=str, separators=(",", ":"))
        digest.update(payload.encode("utf-8"))
        digest.update(b"\n")
    return digest.hexdigest()


def _metrics(headers: list[str], rows: list[tuple[Any, ...]]) -> tuple[dict[str, Any], dict[str, list[int]]]:
    index = {header: position for position, header in enumerate(headers)}
    required = {
        "Код_товару",
        "Назва_позиції",
        "Ціна",
        "Валюта",
        "Наявність",
        "Кількість",
        "Назва_групи",
        "Унікальний_ідентифікатор",
        "Виробник",
        "Продукт_на_сайті",
    }
    missing = sorted(required - set(index))
    if missing:
        raise ValueError("В исходном листе нет колонок: " + ", ".join(missing))

    normalized: defaultdict[str, list[int]] = defaultdict(list)
    raw_oe: defaultdict[str, list[int]] = defaultdict(list)
    stable_skus: list[str] = []
    kemp = 0
    dimensions = 0
    positive_stock = 0
    available_marker = 0
    valid_oe = 0
    urls = 0
    currency = Counter()
    for offset, row in enumerate(rows, start=2):
        raw = _text(row[index["Код_товару"]])
        norm = normalize_identifier(raw)
        normalized[norm].append(offset)
        raw_oe[raw].append(offset)
        stable_skus.append(_text(row[index["Унікальний_ідентифікатор"]]))
        valid_oe += len(norm) >= 3
        brand = _text(row[index["Виробник"]]).upper()
        kemp += brand in {"KEMP", "КЕМР"}
        dimension_headers = ("Вага,кг", "Ширина,см", "Висота,см", "Довжина,см")
        dimensions += all(
            header in index and row[index[header]] not in (None, "")
            for header in dimension_headers
        )
        quantity = row[index["Кількість"]]
        try:
            positive_stock += quantity is not None and float(quantity) > 0
        except (TypeError, ValueError):
            pass
        marker = _text(row[index["Наявність"]])
        available_marker += marker in {"+", "!"}
        urls += bool(_text(row[index["Продукт_на_сайті"]]))
        currency[_text(row[index["Валюта"]]) or "<blank>"] += 1

    collision_groups = {key: value for key, value in normalized.items() if len(value) > 1}
    invalid_rows = [row for key, values in normalized.items() if len(key) < 3 for row in values]
    collision_rows = sorted(row for values in collision_groups.values() for row in values)
    review_rows = sorted(set(invalid_rows) | set(collision_rows))
    raw_duplicate_groups = {key: value for key, value in raw_oe.items() if len(value) > 1}
    duplicate_sku_rows = sum(count - 1 for count in Counter(stable_skus).values() if count > 1)
    total = len(rows)
    auto_admissible = total - len(review_rows)
    cost_headers = [header for header in headers if is_raw_cost_label(header)]
    result = {
        "total_rows": total,
        "valid_rows": auto_admissible,
        "rejected_or_manual_review_rows": len(review_rows),
        "kemp_brand_rows": kemp,
        "kemp_brand_share": _wilson(kemp, total),
        "syntactically_valid_oe_rows": valid_oe,
        "syntactically_valid_oe_share": _wilson(valid_oe, total),
        "distinct_normalized_oe_count": len(normalized),
        "distinct_normalized_oe_per_row": _wilson(len(normalized), total),
        "auto_admissible_oe_rows": auto_admissible,
        "auto_admissible_oe_share": _wilson(auto_admissible, total),
        "dimension_complete_rows": dimensions,
        "dimension_coverage": _wilson(dimensions, total),
        "currency_distribution": dict(sorted(currency.items())),
        "cost_coverage": "sensitive_do_not_export_raw",
        "raw_cost_columns_present": cost_headers,
        "duplicate_raw_oe_groups": len(raw_duplicate_groups),
        "duplicate_raw_oe_rows": sum(len(value) for value in raw_duplicate_groups.values()),
        "normalized_oe_collision_groups": len(collision_groups),
        "normalized_oe_collision_rows": len(collision_rows),
        "duplicate_stable_sku_count": duplicate_sku_rows,
        "missing_product_url_count": total - urls,
        "product_url_coverage": _wilson(urls, total),
        "positive_stock_rows": positive_stock,
        "positive_stock_share": _wilson(positive_stock, total),
        "available_marker_rows": available_marker,
        "available_marker_share": _wilson(available_marker, total),
        "stock_status_coverage": _wilson(0, total),
        "stock_age_days_coverage": _wilson(0, total),
    }
    return result, {
        "review_rows": review_rows,
        "invalid_rows": invalid_rows,
        "collision_rows": collision_rows,
    }


def _style_header(sheet, max_column: int) -> None:
    sheet.row_dimensions[1].height = 34
    for cell in sheet[1][:max_column]:
        cell.font = Font(name="Arial", size=10, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=BORDER)


def _add_validation(sheet, cell_range: str, *, kind: str, formula1: str, formula2: str | None = None, prompt: str) -> None:
    validation = DataValidation(
        type=kind,
        operator="between" if formula2 is not None else None,
        formula1=formula1,
        formula2=formula2,
        allow_blank=True,
        showErrorMessage=True,
        showInputMessage=True,
        errorTitle="Неверное значение",
        error="Проверьте формат ячейки.",
        promptTitle="Ввод Юрия",
        prompt=prompt,
    )
    sheet.add_data_validation(validation)
    validation.add(cell_range)


def _build_input_sheet(workbook, headers: list[str], rows: list[tuple[Any, ...]]) -> None:
    sheet = workbook.create_sheet("Ввод Юрия", 1)
    source_index = {header: position for position, header in enumerate(headers)}
    sheet.append(INPUT_HEADERS)
    manual_start = INPUT_HEADERS.index("Статус товара") + 1
    manual_end = INPUT_HEADERS.index("Комментарий Юрия") + 1
    source_fill = PatternFill("solid", fgColor=PALE_GRAY)
    input_fill = PatternFill("solid", fgColor=PALE_YELLOW)
    cost_fill = PatternFill("solid", fgColor=PALE_ORANGE)
    source_font = Font(name="Arial", size=10, color=TEXT)
    input_font = Font(name="Arial", size=10, color=BLUE)
    link_font = Font(name="Arial", size=10, color=BLUE, underline="single")
    thin_border = Border(bottom=Side(style="hair", color="D8E0E8"))

    for source_row, row in enumerate(rows, start=2):
        marker = _text(row[source_index["Наявність"]])
        quantity = row[source_index["Кількість"]]
        if marker in {"+", "!"}:
            available = True
        elif marker == "-":
            available = False
        else:
            try:
                available = quantity is not None and float(quantity) > 0
            except (TypeError, ValueError):
                available = False
        values = (
            _text(row[source_index["Унікальний_ідентифікатор"]]),
            _text(row[source_index["Код_товару"]]),
            _text(row[source_index["Назва_позиції"]]),
            _text(row[source_index["Назва_групи"]]),
            _text(row[source_index["Виробник"]]),
            float(_text(row[source_index["Ціна"]]).replace(",", ".")),
            _text(row[source_index["Валюта"]]),
            quantity,
            "Да" if available else "Нет",
            _text(row[source_index["Продукт_на_сайті"]]),
            _text(row[source_index.get("Номер_пристрою_(MPN)", -1)]) if "Номер_пристрою_(MPN)" in source_index else "",
            _text(row[source_index["Опис"]]),
            _text(row[source_index["Пошукові_запити"]]),
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            source_row,
            marker,
        )
        sheet.append(values)
        target_row = sheet.max_row
        for column in range(1, len(INPUT_HEADERS) + 1):
            cell = sheet.cell(target_row, column)
            cell.font = input_font if manual_start <= column <= manual_end else source_font
            cell.fill = input_fill if manual_start <= column <= manual_end else source_fill
            cell.border = thin_border
            cell.alignment = Alignment(vertical="top")
        cost_cell = sheet.cell(target_row, INPUT_HEADERS.index("Себестоимость") + 1)
        cost_cell.fill = cost_fill
        url_cell = sheet.cell(target_row, INPUT_HEADERS.index("Ссылка на товар") + 1)
        if url_cell.value:
            url_cell.hyperlink = str(url_cell.value)
            url_cell.font = link_font

    _style_header(sheet, len(INPUT_HEADERS))
    comments = {
        "N1": "Выберите: ходовой, залежалый, неликвид или неизвестно.",
        "O1": "Необязательно. Не выдумывайте число, если точный возраст неизвестен.",
        "P1": "Необязательно. При загрузке сервер примет значение только в режиме SERVER_SIDE_ENCRYPTED.",
        "X1": "Вводите как процент, например 5%.",
        "Y1": "Необязательный множитель ручного приоритета; пусто = 1.",
    }
    for coordinate, text in comments.items():
        sheet[coordinate].comment = Comment(text, "Marko")
    max_row = sheet.max_row
    _add_validation(
        sheet,
        f"N2:N{max_row}",
        kind="list",
        formula1='"ходовой,залежалый,неликвид,неизвестно"',
        prompt="Выберите статус из списка.",
    )
    for column in ("O", "Q", "R", "S", "T", "U", "V", "W"):
        _add_validation(
            sheet,
            f"{column}2:{column}{max_row}",
            kind="decimal",
            formula1="0",
            formula2="1000000000",
            prompt="Оставьте пустым или введите число не меньше 0.",
        )
    _add_validation(
        sheet,
        f"P2:P{max_row}",
        kind="decimal",
        formula1="0.01",
        formula2="999999999999.99",
        prompt="Себестоимость в UAH, больше 0, до 2 знаков.",
    )
    _add_validation(
        sheet,
        f"X2:X{max_row}",
        kind="decimal",
        formula1="0",
        formula2="1",
        prompt="Значение от 0% до 100%.",
    )
    _add_validation(
        sheet,
        f"Y2:Y{max_row}",
        kind="decimal",
        formula1="0.0001",
        formula2="1000000",
        prompt="Положительное число; пусто = 1.",
    )
    for row in range(2, max_row + 1):
        sheet[f"A{row}"].number_format = "@"
        sheet[f"B{row}"].number_format = "@"
        sheet[f"K{row}"].number_format = "@"
        sheet[f"F{row}"].number_format = '#,##0.00;[Red](#,##0.00);-'
        sheet[f"P{row}"].number_format = '#,##0.00;[Red](#,##0.00);-'
        sheet[f"X{row}"].number_format = "0.0%"
    widths = {
        "A": 16,
        "B": 18,
        "C": 48,
        "D": 26,
        "E": 14,
        "F": 12,
        "G": 9,
        "H": 11,
        "I": 11,
        "J": 34,
        "K": 18,
        "L": 48,
        "M": 38,
        "N": 17,
        "O": 15,
        "P": 17,
        "Q": 16,
        "R": 16,
        "S": 16,
        "T": 20,
        "U": 20,
        "V": 18,
        "W": 18,
        "X": 12,
        "Y": 12,
        "Z": 28,
        "AA": 14,
        "AB": 18,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    for column in ("K", "L", "M", "AA", "AB"):
        sheet.column_dimensions[column].hidden = True
    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(INPUT_HEADERS))}{max_row}"
    sheet.sheet_view.showGridLines = False


def _build_review_sheet(workbook, headers: list[str], rows: list[tuple[Any, ...]]) -> None:
    sheet = workbook.create_sheet("Проверить OE", 2)
    index = {header: position for position, header in enumerate(headers)}
    by_norm: defaultdict[str, list[tuple[int, tuple[Any, ...]]]] = defaultdict(list)
    for source_row, row in enumerate(rows, start=2):
        raw = _text(row[index["Код_товару"]])
        by_norm[normalize_identifier(raw)].append((source_row, row))
    review_headers = (
        "Причина",
        "Исходная строка",
        "SKU",
        "OE как в файле",
        "OE normalized",
        "Строк в группе",
        "Название",
        "Категория",
        "Ссылка",
        "Что сделать",
    )
    sheet.append(review_headers)
    for norm, group in sorted(by_norm.items(), key=lambda item: (item[0], item[1][0][0])):
        if len(norm) >= 3 and len(group) == 1:
            continue
        reason = "INVALID_OE" if len(norm) < 3 else "NORMALIZED_OE_COLLISION"
        action = (
            "Исправить OE вручную"
            if reason == "INVALID_OE"
            else "Проверить, один ли это товар или разные модификации"
        )
        for source_row, row in group:
            values = (
                reason,
                source_row,
                _text(row[index["Унікальний_ідентифікатор"]]),
                _text(row[index["Код_товару"]]),
                norm,
                len(group),
                _text(row[index["Назва_позиції"]]),
                _text(row[index["Назва_групи"]]),
                _text(row[index["Продукт_на_сайті"]]),
                action,
            )
            sheet.append(values)
            current = sheet.max_row
            fill = PatternFill("solid", fgColor=PALE_RED if reason == "INVALID_OE" else PALE_YELLOW)
            for cell in sheet[current]:
                cell.font = Font(name="Arial", size=10, color=TEXT)
                cell.fill = fill
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.border = Border(bottom=Side(style="hair", color="D8E0E8"))
            link = sheet.cell(current, 9)
            if link.value:
                link.hyperlink = str(link.value)
                link.font = Font(name="Arial", size=10, color=BLUE, underline="single")
    _style_header(sheet, len(review_headers))
    widths = (25, 17, 16, 22, 22, 17, 52, 28, 35, 48)
    for position, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(position)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:J{sheet.max_row}"
    sheet.sheet_view.showGridLines = False


def _build_instructions(workbook, *, source_hash: str, metrics: dict[str, Any]) -> None:
    sheet = workbook.create_sheet("Инструкция", 0)
    sheet.sheet_view.showGridLines = False
    sheet.merge_cells("A1:H1")
    title = sheet["A1"]
    title.value = "Marko · ручные данные каталога Юрия"
    title.font = Font(name="Arial", size=18, bold=True, color=WHITE)
    title.fill = PatternFill("solid", fgColor=NAVY)
    title.alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 38
    instructions = (
        (3, "1", "Откройте лист «Ввод Юрия». Жёлтые ячейки можно заполнять, серые — это данные Prom.ua."),
        (4, "2", "Для V1 достаточно статуса: ходовой / залежалый / неликвид. Число дней необязательно."),
        (5, "3", "Продажи, просмотры, конверсия и приоритет необязательны. Пустые ячейки не считаются нулями."),
        (6, "4", "Себестоимость необязательна. При загрузке она шифруется на сервере, но в самом XLSX остаётся обычным текстом."),
        (7, "5", "Лист «Проверить OE» содержит только конфликтные OE. Они не допускаются к автоматическому сопоставлению до ручной проверки."),
    )
    for row, number, text in instructions:
        sheet[f"A{row}"] = number
        sheet[f"A{row}"].font = Font(name="Arial", size=12, bold=True, color=WHITE)
        sheet[f"A{row}"].fill = PatternFill("solid", fgColor=BLUE)
        sheet[f"A{row}"].alignment = Alignment(horizontal="center", vertical="center")
        sheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=8)
        cell = sheet.cell(row, 2, text)
        cell.font = Font(name="Arial", size=11, color=TEXT)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.fill = PatternFill("solid", fgColor=PALE_BLUE)
        sheet.row_dimensions[row].height = 46
    sheet["A10"] = "Факты из полученного каталога"
    sheet["A10"].font = Font(name="Arial", size=13, bold=True, color=NAVY)
    facts = (
        ("Товарных строк", metrics["total_rows"]),
        ("Бренд KEMP", f'{metrics["kemp_brand_share"]["point"]:.2%}'),
        ("Различных normalized OE / строк", f'{metrics["distinct_normalized_oe_per_row"]["point"]:.2%}'),
        ("Автодопуск после fail-closed OE-проверки", f'{metrics["valid_rows"]} / {metrics["total_rows"]}'),
        ("Строк для ручной OE-проверки", metrics["rejected_or_manual_review_rows"]),
        ("Полные габариты", f'{metrics["dimension_coverage"]["point"]:.2%}'),
        ("Отмечено доступным в Prom.ua", f'{metrics["available_marker_share"]["point"]:.2%}'),
        ("Положительный остаток", f'{metrics["positive_stock_share"]["point"]:.2%}'),
        ("SHA-256 исходного XLSX", source_hash),
    )
    for offset, (label, value) in enumerate(facts, start=11):
        sheet[f"A{offset}"] = label
        sheet.merge_cells(start_row=offset, start_column=1, end_row=offset, end_column=3)
        sheet[f"A{offset}"].font = Font(name="Arial", size=10, bold=True, color=TEXT)
        sheet[f"D{offset}"] = value
        sheet.merge_cells(start_row=offset, start_column=4, end_row=offset, end_column=8)
        sheet[f"D{offset}"].font = Font(name="Arial", size=10, color=TEXT)
        sheet[f"D{offset}"].alignment = Alignment(wrap_text=True)
        fill = PatternFill("solid", fgColor=PALE_GREEN if offset < 16 else PALE_GRAY)
        sheet[f"A{offset}"].fill = fill
        sheet[f"D{offset}"].fill = fill
    for column in range(1, 9):
        sheet.column_dimensions[get_column_letter(column)].width = 18
    sheet.freeze_panes = "A3"


def build(source: Path, output: Path) -> dict[str, Any]:
    if not source.is_file():
        raise FileNotFoundError(source)
    source_hash = _sha256(source)
    workbook = load_workbook(source, data_only=False)
    if SOURCE_SHEET not in workbook.sheetnames:
        raise ValueError(f"Лист {SOURCE_SHEET!r} не найден")
    original_sheets = list(workbook.sheetnames)
    original_digests = {name: _sheet_digest(workbook[name]) for name in original_sheets}
    for name in GENERATED_SHEETS:
        if name in workbook.sheetnames:
            del workbook[name]
    headers, rows = _source_rows(workbook[SOURCE_SHEET])
    metrics, _review = _metrics(headers, rows)
    _build_instructions(workbook, source_hash=source_hash, metrics=metrics)
    _build_input_sheet(workbook, headers, rows)
    _build_review_sheet(workbook, headers, rows)
    workbook.active = workbook.sheetnames.index("Ввод Юрия")
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    workbook.close()

    verified = load_workbook(output, read_only=True, data_only=False)
    try:
        if verified.active.title != "Ввод Юрия":
            raise AssertionError("Рабочий лист не активен")
        if verified["Ввод Юрия"].max_row != len(rows) + 1:
            raise AssertionError("Число строк рабочего листа изменилось")
        for name, digest in original_digests.items():
            if _sheet_digest(verified[name]) != digest:
                raise AssertionError(f"Исходный лист {name!r} изменён")
        formula_count = sum(
            1
            for name in GENERATED_SHEETS
            for row in verified[name].iter_rows()
            for cell in row
            if cell.data_type == "f"
        )
        if formula_count:
            raise AssertionError(f"Найдено формул: {formula_count}")
    finally:
        verified.close()

    parsed = parse_catalog_xlsx(output.read_bytes())
    if parsed.total_rows != len(rows):
        raise AssertionError("Парсер видит другое число строк")
    if len(parsed.rows) != metrics["valid_rows"]:
        raise AssertionError("Парсер и метрики не согласованы")
    metrics.update(
        {
            "source_path": str(source),
            "source_sha256": source_hash,
            "output_path": str(output),
            "output_sha256": _sha256(output),
            "output_size_bytes": output.stat().st_size,
            "parser_imported_rows": len(parsed.rows),
            "parser_rejected_rows": len(parsed.issues),
            "generated_formula_count": 0,
            "original_sheet_digests_preserved": True,
        }
    )
    metrics_path = output.with_suffix(".metrics.json")
    metrics_path.write_text(
        json.dumps(metrics, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    return metrics


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    print(json.dumps(build(args.source, args.output), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
