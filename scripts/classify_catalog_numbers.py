#!/usr/bin/env python3
"""Classify every catalog identifier for the bounded Metis Task A audit."""

from __future__ import annotations

import argparse
from collections import Counter
import csv
import hashlib
import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from marko.services.catalog_number_nature import (
    CatalogNumberClassification,
    CatalogNumberNature,
    CatalogNumberRecord,
    classify_catalog_numbers,
    load_catalog_number_nature_policy,
    load_live_independent_exact_counts,
)


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SOURCE = ROOT / ".artifacts/yuri_catalog/YURI_CATALOG_MARKO_INPUT_V1.xlsx"
DEFAULT_CONFIG = ROOT / "backend/config/catalog_number_nature.yaml"
DEFAULT_LIVE_OFFERS = (
    ROOT / ".artifacts/metis_next_steps_20260719/METIS_30_OE_OFFERS.csv"
)
DEFAULT_OUTPUT_DIR = ROOT / "outputs/metis_catalog_number_nature_20260724"

FIELD_ALIASES: dict[str, tuple[str, ...]] = {
    "source_row": ("Исходная строка",),
    "sku": ("SKU", "Унікальний_ідентифікатор"),
    "oe_raw": ("OE", "Код_товару"),
    "title": ("Название", "Назва_позиції"),
    "category": ("Категория", "Назва_групи"),
    "brand": ("Бренд", "Виробник"),
    "mpn_raw": ("MPN",),
    "search_queries": ("Поисковые запросы", "Пошукові_запити"),
    "product_url": ("Ссылка на товар", "Продукт_на_сайті"),
}
REQUIRED_FIELDS = frozenset({"sku", "oe_raw", "title"})


def _parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Classify catalog numbers as KEMP internal, manufacturer/universal, "
            "or unknown without inferring fitment."
        )
    )
    parser.add_argument("--input", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--config", type=Path, default=DEFAULT_CONFIG)
    parser.add_argument("--live-offers", type=Path, default=DEFAULT_LIVE_OFFERS)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--sheet-name", default=None)
    return parser.parse_args()


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, float):
        return format(value, ".15g")
    return str(value).strip()


def _resolve_sheet(workbook, requested: str | None):
    if requested:
        if requested not in workbook.sheetnames:
            raise ValueError(f"Workbook does not contain sheet {requested!r}")
        return workbook[requested]
    for candidate in ("Ввод Юрия", "Export Products Sheet"):
        if candidate in workbook.sheetnames:
            return workbook[candidate]
    raise ValueError(
        "Workbook contains neither 'Ввод Юрия' nor 'Export Products Sheet'"
    )


def _column_mapping(headers: list[str]) -> dict[str, int]:
    positions = {header: index for index, header in enumerate(headers)}
    mapping: dict[str, int] = {}
    for field, aliases in FIELD_ALIASES.items():
        position = next(
            (positions[alias] for alias in aliases if alias in positions),
            None,
        )
        if position is not None:
            mapping[field] = position
    missing = sorted(REQUIRED_FIELDS - set(mapping))
    if missing:
        raise ValueError("Workbook is missing required fields: " + ", ".join(missing))
    return mapping


def _read_records(
    source_path: Path,
    *,
    sheet_name: str | None,
) -> tuple[str, list[CatalogNumberRecord]]:
    workbook = load_workbook(source_path, read_only=True, data_only=True)
    try:
        sheet = _resolve_sheet(workbook, sheet_name)
        iterator = sheet.iter_rows(values_only=True)
        headers = [_text(value) for value in next(iterator)]
        mapping = _column_mapping(headers)
        records: list[CatalogNumberRecord] = []
        for workbook_row, row in enumerate(iterator, start=2):
            if not any(value not in (None, "") for value in row):
                continue
            raw_source_row = (
                row[mapping["source_row"]]
                if "source_row" in mapping
                else workbook_row
            )
            try:
                source_row = int(raw_source_row)
            except (TypeError, ValueError):
                source_row = workbook_row
            records.append(
                CatalogNumberRecord(
                    source_row=source_row,
                    sku=_text(row[mapping["sku"]]),
                    oe_raw=_text(row[mapping["oe_raw"]]),
                    title=_text(row[mapping["title"]]),
                    category=_optional_field(row, mapping, "category"),
                    brand=_optional_field(row, mapping, "brand"),
                    mpn_raw=_optional_field(row, mapping, "mpn_raw"),
                    search_queries=_optional_field(row, mapping, "search_queries"),
                    product_url=_optional_field(row, mapping, "product_url"),
                )
            )
        return sheet.title, records
    finally:
        workbook.close()


def _optional_field(
    row: tuple[Any, ...],
    mapping: dict[str, int],
    field: str,
) -> str:
    position = mapping.get(field)
    return _text(row[position]) if position is not None else ""


def _row_payload(
    record: CatalogNumberRecord,
    result: CatalogNumberClassification,
) -> dict[str, Any]:
    return {
        "source_row": record.source_row,
        "sku": record.sku,
        "oe_raw": record.oe_raw,
        "oe_norm": result.oe_norm,
        "number_nature": result.nature.value,
        "reason_code": result.reason_code,
        "rule_id": result.rule_id or "",
        "live_independent_exact_count": result.live_independent_exact_count,
        "detected_kemp_internal_tokens": " | ".join(
            result.detected_internal_tokens
        ),
        "title": record.title,
        "category": record.category,
        "brand": record.brand,
        "mpn_raw": record.mpn_raw,
        "product_url": record.product_url,
    }


def _examples(
    rows: list[dict[str, Any]],
    *,
    nature: CatalogNumberNature,
    limit: int = 20,
) -> list[dict[str, Any]]:
    candidates = [row for row in rows if row["number_nature"] == nature.value]
    reason_order = {
        "LIVE_INDEPENDENT_EXACT": 0,
        "DISTINCT_KEMP_INTERNAL_TOKEN": 1,
        "KEMP_INTERNAL_PATTERN": 1,
        "MANUFACTURER_PATTERN_AND_CONTEXT": 2,
        "CONFLICTING_KEMP_IDENTIFIERS": 1,
        "INVALID_OR_EMPTY_IDENTIFIER": 2,
        "NO_SUFFICIENT_EVIDENCE": 3,
    }
    candidates.sort(
        key=lambda row: (
            reason_order.get(str(row["reason_code"]), 9),
            str(row["rule_id"]),
            int(row["source_row"]),
        )
    )
    return [
        {
            "source_row": row["source_row"],
            "oe_raw": row["oe_raw"],
            "oe_norm": row["oe_norm"],
            "title": row["title"],
            "reason_code": row["reason_code"],
            "rule_id": row["rule_id"] or None,
            "live_independent_exact_count": row[
                "live_independent_exact_count"
            ],
        }
        for row in candidates[:limit]
    ]


def _build_summary(
    *,
    source_path: Path,
    source_sha256: str,
    source_sheet: str,
    config_path: Path,
    live_offers_path: Path | None,
    rows: list[dict[str, Any]],
    live_counts: dict[str, int],
) -> dict[str, Any]:
    total = len(rows)
    nature_counts = Counter(row["number_nature"] for row in rows)
    reason_counts = Counter(row["reason_code"] for row in rows)
    live_target_rows = [row for row in rows if row["oe_norm"] in live_counts]
    live_disagreements = [
        row
        for row in live_target_rows
        if row["number_nature"] != CatalogNumberNature.MANUFACTURER_OE.value
    ]
    return {
        "schema_version": "metis-catalog-number-nature-report-v1",
        "task": "TASK_A_CATALOG_NUMBER_NATURE",
        "source": {
            "catalog_path": str(source_path.resolve()),
            "catalog_sha256": source_sha256,
            "sheet": source_sheet,
            "config_path": str(config_path.resolve()),
            "live_offers_path": (
                str(live_offers_path.resolve()) if live_offers_path else None
            ),
        },
        "totals": {
            "catalog_rows": total,
            "classified_rows": sum(nature_counts.values()),
            "counts_sum_matches_catalog": sum(nature_counts.values()) == total,
        },
        "number_nature": {
            nature.value: {
                "count": nature_counts[nature.value],
                "share": nature_counts[nature.value] / total if total else 0,
            }
            for nature in CatalogNumberNature
        },
        "reason_counts": dict(sorted(reason_counts.items())),
        "live_cross_check": {
            "distinct_numbers_with_independent_exact_evidence": len(live_counts),
            "catalog_rows_with_live_evidence": len(live_target_rows),
            "classified_as_manufacturer_oe": (
                len(live_target_rows) - len(live_disagreements)
            ),
            "disagreements": len(live_disagreements),
        },
        "examples": {
            nature.value: _examples(rows, nature=nature)
            for nature in CatalogNumberNature
        },
        "interpretation": {
            "manufacturer_oe_scope": (
                "A number external to KEMP according to direct market evidence, "
                "a separate KEMP identifier in the same catalog row, or a "
                "configured manufacturer pattern with matching context. This "
                "is not independent OE certification or fitment proof."
            ),
            "unknown_policy": (
                "No direct market evidence, no distinct KEMP token and no "
                "configured manufacturer pattern with matching context."
            ),
            "pricing_impact": (
                "None. Number-nature classification does not admit offers to "
                "calibration or pricing."
            ),
        },
        "gate": {
            "status": (
                "PASS"
                if sum(nature_counts.values()) == total and not live_disagreements
                else "FAIL"
            ),
            "next_stage_authorized": False,
        },
    }


def _write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    if not rows:
        raise ValueError("Cannot write an empty classification table")
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def _write_report(path: Path, summary: dict[str, Any]) -> None:
    natures = summary["number_nature"]
    reason_counts = summary["reason_counts"]
    live = summary["live_cross_check"]
    lines = [
        "# Metis — Task A: природа номеров каталога",
        "",
        "## Итог",
        "",
        "| Тип номера | Количество | Доля |",
        "|---|---:|---:|",
    ]
    labels = {
        CatalogNumberNature.KEMP_INTERNAL.value: "Внутренний номер KEMP",
        CatalogNumberNature.MANUFACTURER_OE.value: (
            "OE / универсальный номер производителя"
        ),
        CatalogNumberNature.UNKNOWN.value: "Неопределённый",
    }
    for nature in CatalogNumberNature:
        metrics = natures[nature.value]
        lines.append(
            f"| {labels[nature.value]} | {metrics['count']:,} | "
            f"{metrics['share']:.2%} |"
        )
    lines.extend(
        [
            "",
            "Классифицировано ровно "
            f"{summary['totals']['classified_rows']:,} из "
            f"{summary['totals']['catalog_rows']:,} строк.",
            "",
            "## Основания",
            "",
            "| Основание | Строк |",
            "|---|---:|",
        ]
    )
    for reason, count in reason_counts.items():
        lines.append(f"| `{reason}` | {count:,} |")
    lines.extend(
        [
            "",
            "## Проверка живыми данными",
            "",
            f"- Независимое exact-evidence есть для {live['distinct_numbers_with_independent_exact_evidence']} номеров.",
            f"- Строк каталога с таким evidence: {live['catalog_rows_with_live_evidence']}.",
            f"- Классифицированы как внешний номер: {live['classified_as_manufacturer_oe']}.",
            f"- Противоречия: {live['disagreements']}.",
            "",
            "## Граница вывода",
            "",
            "Тип `MANUFACTURER_OE` означает внешний относительно KEMP номер: "
            "это подтверждено независимым exact-listing, отдельным KEMP-кодом "
            "в той же строке каталога либо консервативным manufacturer-pattern "
            "с контекстом марки. Каталожная связь не является независимой OE-"
            "сертификацией. Класс не доказывает применимость объявления и "
            "ничего не допускает в calibration/pricing.",
            "",
            f"**TASK_A_GATE: {summary['gate']['status']}**  ",
            "**NEXT_STAGE_AUTHORIZED: false**",
            "",
        ]
    )
    path.write_text("\n".join(lines), encoding="utf-8")


def main() -> None:
    args = _parse_args()
    source_path = args.input.expanduser().resolve()
    config_path = args.config.expanduser().resolve()
    live_path = (
        args.live_offers.expanduser().resolve() if args.live_offers else None
    )
    output_dir = args.output_dir.expanduser().resolve()
    output_dir.mkdir(parents=True, exist_ok=True)

    policy = load_catalog_number_nature_policy(config_path)
    source_sha256 = _sha256(source_path)
    if source_sha256 != policy.source_catalog_sha256:
        raise ValueError(
            "Catalog SHA-256 does not match the client-specific policy: "
            f"{source_sha256} != {policy.source_catalog_sha256}"
        )

    source_sheet, records = _read_records(source_path, sheet_name=args.sheet_name)
    live_counts_mapping = load_live_independent_exact_counts(
        live_path,
        owned_seller_ids=policy.owned_seller_ids,
    )
    results = classify_catalog_numbers(
        records,
        policy=policy,
        live_independent_exact_counts=live_counts_mapping,
    )
    if len(results) != len(records):
        raise RuntimeError("Classifier did not return exactly one result per row")

    output_rows = [
        _row_payload(record, result)
        for record, result in zip(records, results, strict=True)
    ]
    summary = _build_summary(
        source_path=source_path,
        source_sha256=source_sha256,
        source_sheet=source_sheet,
        config_path=config_path,
        live_offers_path=live_path,
        rows=output_rows,
        live_counts=dict(live_counts_mapping),
    )

    csv_path = output_dir / "METIS_CATALOG_NUMBER_NATURE_ROWS.csv"
    summary_path = output_dir / "METIS_CATALOG_NUMBER_NATURE_SUMMARY.json"
    report_path = output_dir / "METIS_CATALOG_NUMBER_NATURE_REPORT.md"
    _write_csv(csv_path, output_rows)
    summary_path.write_text(
        json.dumps(summary, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    _write_report(report_path, summary)

    print(
        json.dumps(
            {
                "catalog_rows": len(records),
                "number_nature": summary["number_nature"],
                "live_cross_check": summary["live_cross_check"],
                "gate": summary["gate"],
                "outputs": {
                    "rows_csv": str(csv_path),
                    "summary_json": str(summary_path),
                    "report_md": str(report_path),
                },
            },
            ensure_ascii=False,
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
