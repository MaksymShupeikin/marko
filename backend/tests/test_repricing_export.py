"""The .xlsx a run produces: metadata in the header, hidden rows left out."""
from __future__ import annotations

from datetime import UTC, datetime
from decimal import Decimal
from types import SimpleNamespace
from uuid import uuid4

from openpyxl import load_workbook

from marko.infrastructure.db.models import (
    RepriceItemStatus,
    RepriceMode,
    RepriceOutcome,
    RepricePolicy,
    RepriceRun,
    RepriceRunItem,
    RepriceScope,
)
from marko.services import repricing_export


def _run() -> RepriceRun:
    run = RepriceRun(
        id=uuid4(),
        workspace_id=uuid4(),
        scope=RepriceScope.partial,
        mode=RepriceMode.resume,
        policy=RepricePolicy.balanced,
        engine="legacy_min_minus",
        catalog_scope_signature="abc123",
        catalog_item_count=4901,
        store_ids=["store-1"],
        changed_count=2,
        unchanged_count=1,
        skipped_count=1,
        failed_count=0,
    )
    run.created_at = datetime(2026, 9, 11, 14, 30, tzinfo=UTC)
    return run


def _row(
    *,
    sku: str,
    old: Decimal | None,
    new: Decimal | None,
    outcome: RepriceOutcome,
    reason: str | None = None,
    price_at_compute: Decimal | None = None,
    current_price: Decimal | None = None,
):
    item = RepriceRunItem(
        run_id=uuid4(),
        listing_id=uuid4(),
        position=0,
        status=RepriceItemStatus.done,
        outcome=outcome,
        reason=reason,
        old_price=old,
        new_price=new,
        delta_abs=(new - old) if (old is not None and new is not None) else None,
        price_at_compute=price_at_compute,
        offers_total=12,
    )
    listing = SimpleNamespace(
        sku=sku,
        name=f"Товар {sku}",
        brand="KEMP",
        currency="UAH",
        url=f"https://prom.ua/{sku}.html",
        current_price=current_price if current_price is not None else old,
    )
    store = SimpleNamespace(name="Автозапчастини")
    return (item, listing, store, None)


def _sheet(rows):
    buffer = repricing_export.build_workbook(_run(), rows)
    return load_workbook(buffer).active


def test_header_records_when_and_against_which_catalog():
    """Через тиждень файл має сам відповідати, коли і по чому його зробили."""
    sheet = _sheet([_row(sku="A1", old=Decimal("100"), new=Decimal("90"),
                         outcome=RepriceOutcome.changed)])
    header = {row[0]: row[1] for row in sheet.iter_rows(values_only=True) if row[0]}

    assert header["Дата прогону"] == "2026-09-11 14:30"
    assert header["Каталог (підпис складу)"] == "abc123"
    assert header["Товарів у каталозі на той момент"] == 4901
    assert "частина каталогу" in header["Охоплення"]
    assert "продовження" in header["Охоплення"]


def test_every_outcome_keeps_its_own_label_and_reason():
    """«Без змін» і «не пораховано» — різні відповіді, змішувати їх не можна."""
    sheet = _sheet(
        [
            _row(sku="A1", old=Decimal("100"), new=Decimal("90"),
                 outcome=RepriceOutcome.changed),
            _row(sku="A2", old=Decimal("100"), new=Decimal("100"),
                 outcome=RepriceOutcome.unchanged),
            _row(sku="A3", old=Decimal("100"), new=None,
                 outcome=RepriceOutcome.no_recommendation,
                 reason="Тонкий ринок"),
        ]
    )
    rows = [row for row in sheet.iter_rows(values_only=True) if row[0] in {"A1", "A2", "A3"}]
    by_sku = {row[0]: row for row in rows}

    assert by_sku["A1"][9] == "Змінено"
    assert by_sku["A2"][9] == "Без змін"
    assert by_sku["A3"][9] == "Не пораховано"
    assert by_sku["A3"][10] == "Тонкий ринок"
    assert by_sku["A3"][5] is None


def test_prices_are_numbers_not_text():
    """Інакше в Excel не порахувати ні суму, ні сортування за різницею."""
    sheet = _sheet([_row(sku="A1", old=Decimal("1950"), new=Decimal("1780"),
                         outcome=RepriceOutcome.changed)])
    row = next(r for r in sheet.iter_rows(values_only=True) if r[0] == "A1")

    assert row[4] == 1950
    assert row[5] == 1780
    assert row[6] == -170
    # Головне — що це число, а не рядок: інакше Excel не порахує ні суму,
    # ні сортування. Цілі значення openpyxl віддає як int, і це нормально.
    assert not isinstance(row[4], str)
    assert not isinstance(row[6], str)


def test_price_drift_since_the_calculation_is_flagged():
    sheet = _sheet(
        [
            _row(sku="A1", old=Decimal("100"), new=Decimal("90"),
                 outcome=RepriceOutcome.changed,
                 price_at_compute=Decimal("100"), current_price=Decimal("120")),
            _row(sku="A2", old=Decimal("100"), new=Decimal("90"),
                 outcome=RepriceOutcome.changed,
                 price_at_compute=Decimal("100"), current_price=Decimal("100")),
        ]
    )
    rows = {r[0]: r for r in sheet.iter_rows(values_only=True) if r[0] in {"A1", "A2"}}
    assert rows["A1"][15] == "так"
    assert rows["A2"][15] == "ні"


def test_column_header_matches_the_row_width():
    sheet = _sheet([_row(sku="A1", old=Decimal("100"), new=Decimal("90"),
                         outcome=RepriceOutcome.changed)])
    values = [row for row in sheet.iter_rows(values_only=True)]
    header = next(row for row in values if row[0] == "Артикул")
    body = next(row for row in values if row[0] == "A1")

    assert len([cell for cell in header if cell is not None]) == len(
        repricing_export._COLUMNS
    )
    assert len(body) == len(header)
