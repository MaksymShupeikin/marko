"""The OE workbook is what the customer imports, so its OE column is a contract.

Two things can happen here that nothing downstream can undo. An internal ``776``
shelf code in the OE column travels to Prom as an original number and matches
nothing; a number the catalogue never really asserted travels as if it were
confirmed. The first kills the export, the second is routed to a sheet a human
reads. Neither is ever silently written into the import sheet.
"""

from __future__ import annotations

from decimal import Decimal

import pytest
from openpyxl import Workbook

from marko.services.oe_catalog_export import (
    IMPORT_HEADERS,
    candidate_numbers,
    oe_named_by_an_asserting_source,
    rests_only_on_our_own_label,
    ExportRow,
    OeCatalogExportError,
    import_values,
    oe_needs_a_human,
    resolve_import_row,
)
from marko.services.xlsx_catalog import (
    FIELD_ALIASES,
    _normalize_header,
    parse_catalog_xlsx,
)


BRANDS = frozenset({"pierburg", "bremi", "febi", "sachs"})


def _row(oe: str, **overrides: str) -> ExportRow:
    fields = {
        "sku": "1153724202",
        "oe": oe,
        "name": "Радиатор Iveco 625*440",
        "category": "7,3 Радиаторы",
        "price": "3297",
        "currency": "UAH",
        "brand": "KEMP",
        "mpn": "93818439",
        "available": "+",
        "stock": "5",
        "url": "https://kemp-cs2847093.prom.ua/p1153724202-radiator.html",
        "internal_code": "7764321",
        "sources": "KEMP_REFERENCE_MAP_V2",
        "other_numbers": "",
        "anomalies": "",
        "evidence_url": "",
        "candidates": "",
        "no_oe_reason": "",
    }
    fields.update(overrides)
    return ExportRow(**fields)


def test_an_internal_code_in_the_oe_column_stops_the_export() -> None:
    with pytest.raises(OeCatalogExportError) as error:
        resolve_import_row(_row("7764321"), brands=BRANDS)

    assert "7764321" in str(error.value)


def test_a_row_without_a_confirmed_number_travels_with_an_empty_oe() -> None:
    """Every priced position is in the table; only the OE cell may be empty."""

    resolved = resolve_import_row(
        _row("", no_oe_reason="оригинальный номер не назван ни одним источником"),
        brands=BRANDS,
    )

    assert resolved.oe == ""
    assert resolved.no_oe_reason


def test_a_plain_oem_number_keeps_its_place_in_the_oe_column() -> None:
    resolved = resolve_import_row(_row("058133753D"), brands=BRANDS)

    assert resolved.oe == "058133753D"
    assert resolved.no_oe_reason == ""


def test_a_number_too_short_to_be_an_oe_moves_out_of_the_oe_column() -> None:
    resolved = resolve_import_row(_row("KL228"), brands=BRANDS)

    assert resolved.oe == ""
    assert "KL228" in resolved.candidates
    assert "коротк" in resolved.no_oe_reason.lower()


def test_a_brand_glued_to_the_tail_moves_out_of_the_oe_column() -> None:
    resolved = resolve_import_row(_row("191906090PIERBURG"), brands=BRANDS)

    assert resolved.oe == ""
    assert "PIERBURG" in resolved.no_oe_reason
    assert "191906090PIERBURG" in resolved.candidates


def test_a_number_without_a_digit_moves_out_of_the_oe_column() -> None:
    resolved = resolve_import_row(_row("ABCDEFGH"), brands=BRANDS)

    assert resolved.oe == ""
    assert resolved.no_oe_reason


def test_the_reason_is_none_only_when_nothing_is_wrong() -> None:
    assert oe_needs_a_human("058133753D", brands=BRANDS) is None
    assert oe_needs_a_human("KL228", brands=BRANDS) is not None


def test_every_import_column_claims_at_most_one_catalog_field() -> None:
    """Two columns claiming one field is the importer's ambiguity error."""

    claimed: dict[str, list[str]] = {}
    for header in IMPORT_HEADERS:
        normalized = _normalize_header(header)
        for field, aliases in FIELD_ALIASES.items():
            if any(normalized == _normalize_header(alias) for alias in aliases):
                claimed.setdefault(field, []).append(header)

    assert not [
        (field, headers) for field, headers in claimed.items() if len(headers) > 1
    ]


def test_the_import_sheet_carries_the_fields_the_importer_requires() -> None:
    """name, category and price are required; sku, oe and url are the point."""

    claimed = {
        field
        for field, aliases in FIELD_ALIASES.items()
        for header in IMPORT_HEADERS
        if any(_normalize_header(header) == _normalize_header(alias) for alias in aliases)
    }

    assert {"sku", "oe", "name", "category", "price", "product_url"} <= claimed


def test_the_import_sheet_reads_back_through_the_customer_importer(tmp_path) -> None:
    """The whole point of the file: the UI must accept it without a mapping.

    The first build shipped the stock *freshness* word in the availability
    column and the importer refused all 3283 rows — it parsed the headers fine
    and rejected every value. Headers matching is not the contract; a row
    surviving the parser is.
    """

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(list(IMPORT_HEADERS))
    sheet.append(import_values(_row("058133753D", available="+")))
    sheet.append(import_values(_row("357905851D", sku="1153724214", available="")))
    sheet.append(import_values(_row("", sku="1153724215", available="+")))
    path = tmp_path / "oe.xlsx"
    workbook.save(path)

    parsed = parse_catalog_xlsx(path.read_bytes())

    assert parsed.issues == []
    assert [row.oe_raw for row in parsed.rows] == ["058133753D", "357905851D", ""]
    assert parsed.column_mapping["oe"] == "OE номер"
    assert parsed.column_mapping["product_url"] == "Ссылка"
    assert parsed.rows[0].current_price == Decimal("3297")
    assert parsed.rows[0].is_available is True
    assert parsed.rows[1].is_available is None
    assert parsed.rows[2].oe_norm == ""


def test_a_row_the_importer_would_refuse_is_named_as_such() -> None:
    """The row still travels — the customer asked for the whole catalogue."""

    resolved = resolve_import_row(_row("058133753D", price=""), brands=BRANDS)

    assert resolved.oe == "058133753D"
    assert "цен" in resolved.no_oe_reason.lower()


def test_our_shelf_code_is_not_offered_even_as_a_candidate() -> None:
    """Nobody outside this company can search for a 776 number."""

    assert candidate_numbers(["776416", "056121113D", "050121113C"]) == (
        "050121113C, 056121113D"
    )
    assert candidate_numbers(["056121113D", "056121113D", ""]) == "056121113D"


ASSERTING = frozenset({"KEMP_REFERENCE_MAP_V2", "SPARETO_OE_PAGE"})


def test_the_anchor_wins_when_a_source_claiming_oe_named_it() -> None:
    picked = oe_named_by_an_asserting_source(
        [("058133753D", ("KEMP_REFERENCE_MAP_V2",)), ("X1", ("KEMP_SITE",))],
        asserting=ASSERTING,
    )

    assert picked == ("058133753D", ("KEMP_REFERENCE_MAP_V2",))


def test_a_supplier_article_anchor_yields_to_the_number_the_book_named() -> None:
    """776457 shipped as AUTC7163 while the book and spareto named 441407151C."""

    picked = oe_named_by_an_asserting_source(
        [
            ("AUTC7163", ("OWN_EXPORT_CODE", "KEMP_REFERENCE_ARTICLE")),
            ("441407151C", ("KEMP_REFERENCE_MAP_V2", "SPARETO_OE_PAGE")),
        ],
        asserting=ASSERTING,
    )

    assert picked == ("441407151C", ("KEMP_REFERENCE_MAP_V2", "SPARETO_OE_PAGE"))


def test_no_number_claimed_as_an_oe_means_no_number_at_all() -> None:
    assert (
        oe_named_by_an_asserting_source(
            [("AUTC7163", ("OWN_EXPORT_CODE", "KEMP_REFERENCE_ARTICLE"))],
            asserting=ASSERTING,
        )
        is None
    )


def test_a_number_only_our_own_card_calls_original_is_not_a_price_key() -> None:
    """The config says so itself: kemp.ua field labels are unreliable (NO_9)."""

    asserting = {"KEMP_SITE", "KEMP_REFERENCE_MAP_V2"}

    assert rests_only_on_our_own_label(
        ("KEMP_REFERENCE_ARTICLE", "KEMP_SITE"), asserting=asserting
    )
    assert not rests_only_on_our_own_label(
        ("KEMP_REFERENCE_MAP_V2", "KEMP_SITE"), asserting=asserting
    )


def test_the_importer_reads_the_internal_code_column(tmp_path) -> None:
    """Наш код KEMP — это то, чем позиция связывается с напарсенной витриной.

    До этого он доезжал до базы только внутри сырой строки: колонки, по которой
    можно искать и соединять, не существовало, и связь каталог ↔ витрина шла по
    бренду с артикулом. 3767 строк из 4901 несут этот код и ждали, чтобы его
    прочитали.
    """

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(list(IMPORT_HEADERS))
    sheet.append(import_values(_row("058133753D")))
    path = tmp_path / "oe.xlsx"
    workbook.save(path)

    parsed = parse_catalog_xlsx(path.read_bytes())

    assert parsed.column_mapping["internal_code"] == "Внутренний код"
    assert parsed.rows[0].internal_code_raw == "7764321"
    assert parsed.rows[0].internal_code_norm == "7764321"


def test_a_number_that_is_not_our_shelf_code_is_refused_as_one(tmp_path) -> None:
    """В эту колонку принимается только внутренний код, а не любой номер.

    Иначе связь каталога с витриной пойдёт по чужому номеру и соединит разные
    позиции.
    """

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(list(IMPORT_HEADERS))
    sheet.append(import_values(_row("058133753D", internal_code="058133753D")))
    path = tmp_path / "oe.xlsx"
    workbook.save(path)

    parsed = parse_catalog_xlsx(path.read_bytes())

    assert parsed.rows[0].internal_code_norm == ""
    assert any("внутренн" in issue.message.lower() for issue in parsed.issues)
