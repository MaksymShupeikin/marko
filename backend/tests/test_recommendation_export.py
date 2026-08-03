from __future__ import annotations

import csv
from datetime import UTC, datetime
from decimal import Decimal
from io import BytesIO, StringIO
from types import SimpleNamespace
from uuid import uuid4

from openpyxl import load_workbook
import pytest

from marko.services import recommendation_export
from marko.services.recommendation_export import (
    RecommendationExportError,
    _export_row,
    _format_money,
    _to_csv,
    _to_xlsx,
)


def _row():
    recommendation = SimpleNamespace(
        action="RAISE",
        current_price=Decimal("100.004"),
        fair_price=Decimal("110.004"),
        recommended_price=Decimal("105.005"),
        absolute_recommended_change=Decimal("5.005"),
        percentage_recommended_change=Decimal("0.05005"),
        confidence=Decimal("0.9"),
        confidence_grade="A",
        reason_codes=["MARKET_SUPPORTS_RAISE"],
        calculation_trace={},
        currency="UAH",
        price_tick=Decimal("0.01"),
        price_tick_version="tick-v1",
        computed_at=datetime(2026, 7, 30, tzinfo=UTC),
    )
    item = SimpleNamespace(
        sku="SKU-1",
        oe_norm="OE-1",
        name="Part",
        category="Filters",
    )
    return _export_row(
        recommendation,
        item,
        source_urls=(
            "https://prom.ua/ua/p1-part.html",
            "https://prom.ua/ua/p2-part.html",
        ),
    )


def test_money_format_matches_ui_tick_rounding_contract() -> None:
    assert (
        _format_money(
            Decimal("105.005"),
            currency="uah",
            price_tick=Decimal("0.01"),
        )
        == "105.01 UAH"
    )
    assert (
        _format_money(
            Decimal("105.4"),
            currency="UAH",
            price_tick=Decimal("1.0000"),
            price_tick_version="uah-integer-v1",
        )
        == "105 UAH"
    )


def test_csv_export_contains_active_rows_and_urls_but_never_cost() -> None:
    payload = _to_csv([_row()]).decode("utf-8-sig")
    parsed = list(csv.DictReader(StringIO(payload)))

    assert len(parsed) == 1
    assert parsed[0]["recommended_price"] == "105.01 UAH"
    assert "https://prom.ua/ua/p1-part.html" in parsed[0]["source_urls"]
    assert all("cost" not in header.casefold() for header in parsed[0])
    assert "себесто" not in payload.casefold()


def test_xlsx_export_has_clickable_source_columns_and_no_cost() -> None:
    workbook = load_workbook(BytesIO(_to_xlsx([_row()])), read_only=False)
    try:
        sheet = workbook["Recommendations"]
        headers = [cell.value for cell in sheet[1]]
        first_url_column = headers.index("source_url_1") + 1
        second_url_column = headers.index("source_url_2") + 1

        assert all("cost" not in str(header).casefold() for header in headers)
        assert sheet.cell(2, first_url_column).hyperlink.target.endswith(
            "/p1-part.html"
        )
        assert sheet.cell(2, second_url_column).hyperlink.target.endswith(
            "/p2-part.html"
        )
    finally:
        workbook.close()


def test_export_exposes_gated_customer_target_without_calling_it_automatic() -> None:
    recommendation = SimpleNamespace(
        action="MANUAL_REVIEW",
        current_price=Decimal("1200"),
        fair_price=Decimal("1000"),
        recommended_price=None,
        absolute_recommended_change=None,
        percentage_recommended_change=None,
        confidence=Decimal("0.8"),
        confidence_grade="MANUAL",
        reason_codes=["COMPARABILITY_AUTOMATIC_ACTIVATION_BLOCKED"],
        calculation_trace={
            "advisory_decision": {
                "action": "LOWER",
                "recommended_price": "980",
                "target_band_low": "950",
                "target_band_high": "980",
                "automatic_price_application": False,
            }
        },
        currency="UAH",
        price_tick=Decimal("1"),
        price_tick_version="uah-integer-v1",
        computed_at=datetime(2026, 7, 30, tzinfo=UTC),
    )
    item = SimpleNamespace(
        sku="SKU-1",
        oe_norm="OE-1",
        name="Part",
        category="Filters",
    )

    row = _export_row(recommendation, item, source_urls=())

    assert row["action"] == "MANUAL_REVIEW"
    assert row["recommended_price"] == ""
    assert row["customer_advisory_action"] == "LOWER"
    assert row["customer_advisory_price"] == "980 UAH"
    assert row["customer_target_band_low"] == "950 UAH"
    assert row["customer_target_band_high"] == "980 UAH"
    assert row["automatic_price_application"] == "false"


@pytest.mark.asyncio
async def test_export_refuses_more_than_the_documented_limit(monkeypatch) -> None:
    async def too_many(*_args, **_kwargs):
        return [], 5001, uuid4(), {
            "raise": 0,
            "lower": 0,
            "review": 0,
            "hold": 0,
            "total": 5001,
        }

    monkeypatch.setattr(recommendation_export, "list_recommendations", too_many)

    with pytest.raises(RecommendationExportError) as error:
        await recommendation_export.export_recommendations(
            SimpleNamespace(),
            workspace_id=uuid4(),
            export_format="csv",
            run_id=None,
            action=None,
            confidence_grade=None,
            category=None,
            queue="all",
            priority_score_type=None,
            confidence_min=None,
            confidence_max=None,
            sort="NEWEST",
        )

    assert error.value.code == "RECOMMENDATION_EXPORT_LIMIT_EXCEEDED"
