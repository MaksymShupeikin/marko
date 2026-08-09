"""Validated XLSX ingestion for Prom catalog exports.

The parser keeps the original row next to normalized pricing inputs.  It does
not execute workbook formulas and never converts article numbers to integers,
so leading zeroes remain intact when they are present in the workbook.
"""

from __future__ import annotations

import asyncio
from collections.abc import Mapping, Sequence
import hashlib
import json
import re
from dataclasses import asdict, dataclass, replace
from datetime import UTC, date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any
from uuid import UUID
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy import func, select, text
from sqlalchemy.ext.asyncio import AsyncSession

from metis.identifiers import normalize_oem_identifier

from marko.core.config import Settings, backend_config_path, get_settings
from marko.infrastructure.db.models import (
    CatalogImportBatch,
    CatalogItem,
    PricingRecommendation,
    PricingRun,
    PricingRunItem,
)
from marko.services.catalog_characteristics import (
    CharacteristicsConfig,
    CharacteristicsExtraction,
    characteristic_column_pairs,
    collect_characteristics,
    extract_characteristics,
    load_characteristics_config,
    normalize_characteristic_name,
)
from marko.services.catalog_costs import add_encrypted_cost_record
from marko.services.catalog_identity_safety import is_internal_catalog_code
from marko.services.cost_privacy import (
    CostPrivacyMode,
    is_raw_cost_label,
    require_server_cost_input_allowed,
)
from marko.services.pricing_runs import (
    IDENTITY_BLOCKED_RECOMMENDATION_ACTION,
    recommendation_price_identity_allowed,
)
from marko.services.unified_catalog import upsert_xlsx_products

MAX_XLSX_BYTES = 25 * 1024 * 1024
MAX_UNCOMPRESSED_XLSX_BYTES = 250 * 1024 * 1024
MAX_ROWS = 100_000
MAX_COLUMNS = 256
MAX_ERROR_LOG = 2_000
CATALOG_ROW_OUTCOMES_CONTRACT_VERSION = "catalog-row-outcomes-v1"

_HEADER_CLEAN_RE = re.compile(r"[^a-zа-яёіїґєԁөү0-9]+", re.IGNORECASE)
_IDENTIFIER_SPLIT_RE = re.compile(r"[,;|\n\r]+")

FIELD_ALIASES: dict[str, tuple[str, ...]] = {
    "sku": (
        "sku",
        "артикул",
        "артикул продавца",
        "артикул продавця",
        "код товара",
        "код товару",
        "код позиции",
    ),
    "oe": (
        "oe",
        "oem",
        "oe oem",
        "oem номер",
        "oe номер",
        "номер oe",
        "номер oem",
        "ое номер",
        "оригинальный номер",
        "оригінальний номер",
        "номер запчасти",
        "номер запчастини",
    ),
    "mpn": ("mpn", "manufacturer part number", "номер производителя"),
    # Наш собственный код позиции (776…). По нему каталог связывается с тем,
    # что напарсено с витрины: у карточки Prom он напечатан в характеристике
    # «Код запчастини», у нас лежит в справочнике. Ни артикул продавца, ни OE
    # этой роли не выполняют — первый принадлежит площадке, второй детали.
    "internal_code": (
        "внутренний код",
        "внутрішній код",
        "код kemp",
        "код кемп",
        "код запчасти",
        "код запчастини",
    ),
    "name": (
        "name",
        "title",
        "название",
        "назва",
        "наименование",
        "название позиции",
        "назва позиції",
    ),
    "category": (
        "category",
        "категория",
        "категорія",
        "группа",
        "група",
        "название группы",
        "назва групи",
    ),
    "price": ("price", "цена", "ціна", "цена розничная", "ціна роздрібна"),
    "currency": ("currency", "валюта"),
    "available": (
        "available",
        "availability",
        "наличие",
        "наявність",
        "статус наличия",
    ),
    "brand": ("brand", "manufacturer", "бренд", "производитель", "виробник"),
    "description": ("description", "описание", "опис"),
    "product_url": (
        "url",
        "product url",
        "ссылка",
        "посилання",
        "ссылка на товар",
        "посилання на товар",
    ),
    "stock_status": (
        "stock status",
        "статус товара",
        "статус товару",
        "статус запасаса",
    ),
    "stock_qty": (
        "stock qty",
        "quantity",
        "остаток",
        "залишок",
        "количество",
        "кількість",
    ),
    "stock_age_days": (
        "stock age days",
        "age days",
        "возраст запаса",
        "вік запасу",
        "дней на складе",
    ),
    "expected_units_sold": (
        "expected units sold",
        "monthly sales",
        "продажи в месяц",
        "продажі на місяць",
        "план продаж в месяц",
        "план продажів на місяць",
    ),
    "units_sold_30d": (
        "units sold 30d",
        "sales 30d",
        "продажи за 30 дней",
        "продажі за 30 днів",
    ),
    "units_sold_60d": (
        "units sold 60d",
        "sales 60d",
        "продажи за 60 дней",
        "продажі за 60 днів",
    ),
    "units_sold_90d": (
        "units sold 90d",
        "sales 90d",
        "продажи за 90 дней",
        "продажі за 90 днів",
    ),
    "days_since_last_sale": (
        "days since last sale",
        "дней с последней продажи",
        "днів з останнього продажу",
    ),
    "historical_monthly_units": (
        "historical monthly units",
        "historical monthly sales",
        "средние продажи в месяц",
        "середні продажі на місяць",
    ),
    "views_30d": (
        "views 30d",
        "просмотры за 30 дней",
        "перегляди за 30 днів",
    ),
    "conversion_rate_proxy": (
        "conversion rate proxy",
        "conversion",
        "конверсия",
        "конверсія",
    ),
    # Raw cost is never ingested through a generic workbook boundary.  The
    # privacy mode must be decided before a dedicated path can exist.
    "cost": (),
    "manual_priority": ("manual priority", "приоритет", "пріоритет"),
}

REQUIRED_FIELDS = frozenset({"name", "category", "price"})


class CatalogImportError(ValueError):
    """The workbook as a whole cannot be imported."""


class SensitiveCatalogImportBlocked(CatalogImportError):
    """A workbook contains raw cost but encrypted ingestion is unavailable."""


@dataclass(frozen=True)
class ImportIssue:
    row: int
    code: str
    message: str


@dataclass(frozen=True)
class ParsedCatalogRow:
    source_row: int
    sku: str
    oe_raw: str
    oe_norm: str
    mpn_raw: str
    mpn_norm: str
    #: Наш внутренний код позиции, как он написан в файле и после нормализации.
    #: ``internal_code_norm`` пуст, если в колонке оказался не наш код: связывать
    #: каталог с витриной по чужому номеру — значит соединить разные позиции.
    internal_code_raw: str
    internal_code_norm: str
    name: str
    category: str
    brand: str | None
    description: str | None
    product_url: str | None
    current_price: Decimal
    currency: str
    is_available: bool | None
    stock_status: str
    stock_qty: Decimal | None
    stock_age_days: Decimal | None
    expected_units_sold: Decimal | None
    units_sold_30d: Decimal | None
    units_sold_60d: Decimal | None
    units_sold_90d: Decimal | None
    days_since_last_sale: Decimal | None
    historical_monthly_units: Decimal | None
    views_30d: Decimal | None
    conversion_rate_proxy: Decimal | None
    cost: Decimal | None
    manual_priority: Decimal
    raw_row: dict[str, Any]
    part_numbers_raw: list[str]
    part_numbers_norm: list[str]
    applicability_brands: list[str]
    applicability_models: list[str]
    characteristics_raw: dict[str, list[str]]
    identity_status: str
    identity_reason: str | None


@dataclass(frozen=True)
class ParsedCatalog:
    rows: list[ParsedCatalogRow]
    issues: list[ImportIssue]
    column_mapping: dict[str, str]
    total_rows: int
    sensitive_costs: dict[int, Decimal]
    characteristics_report: dict[str, Any]


def _canonical_json_sha256(value: object) -> str:
    encoded = json.dumps(
        value,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    ).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


def catalog_row_outcomes(parsed: ParsedCatalog) -> list[dict[str, Any]]:
    """Return one immutable terminal result for every non-empty source row."""

    accepted = {row.source_row: row for row in parsed.rows}
    issues: dict[int, list[ImportIssue]] = {}
    for issue in parsed.issues:
        issues.setdefault(issue.row, []).append(issue)
    overlap = sorted(set(accepted) & set(issues))
    if overlap:
        raise CatalogImportError(
            f"Catalog row accounting overlap for source rows {overlap[:20]}"
        )
    source_rows = sorted(set(accepted) | set(issues))
    if len(source_rows) != parsed.total_rows:
        raise CatalogImportError(
            "Catalog row accounting mismatch: "
            f"terminal={len(source_rows)}, expected={parsed.total_rows}"
        )
    outcomes: list[dict[str, Any]] = []
    for ordinal, source_row in enumerate(source_rows, start=1):
        accepted_row = accepted.get(source_row)
        row_issues = issues.get(source_row, [])
        identity_missing = bool(
            accepted_row is not None
            and accepted_row.identity_status == "UNRESOLVED"
        )
        reason_codes = list(dict.fromkeys(issue.code for issue in row_issues))
        if identity_missing:
            reason_codes.append("CUSTOMER_IDENTITY_MISSING")
        if accepted_row is not None and accepted_row.identity_reason == (
            "NORMALIZED_OE_COLLISION"
        ):
            reason_codes.append("NORMALIZED_OE_COLLISION")
        outcomes.append(
            {
                "source_ordinal": ordinal,
                "source_row": source_row,
                "terminal_status": (
                    "IMPORTED"
                    if accepted_row is not None
                    else "REJECTED_NOT_IMPORTABLE"
                ),
                "sku": accepted_row.sku if accepted_row is not None else None,
                "oe_norm": accepted_row.oe_norm if accepted_row is not None else None,
                "identity_status": (
                    accepted_row.identity_status
                    if accepted_row is not None
                    else None
                ),
                "matching_terminal_status": (
                    "NOT_MATCHED_IDENTITY_COLLISION"
                    if identity_missing
                    and accepted_row is not None
                    and accepted_row.identity_reason == "NORMALIZED_OE_COLLISION"
                    else "NOT_MATCHED_CUSTOMER_IDENTITY_MISSING"
                    if identity_missing
                    else "ELIGIBLE_FOR_IDENTITY_REVIEW"
                    if accepted_row is not None
                    else "NOT_APPLICABLE_NOT_IMPORTED"
                ),
                "reason_codes": reason_codes,
                "details": [issue.message[:1000] for issue in row_issues],
            }
        )
    return outcomes


@dataclass(frozen=True)
class CatalogSheetPreview:
    name: str
    row_count: int
    headers: list[str]
    suggested_mapping: dict[str, str]
    mapping_error: str | None
    sample_rows: list[dict[str, Any]]
    is_catalog_candidate: bool


@dataclass(frozen=True)
class CatalogWorkbookPreview:
    filename: str
    content_sha256: str
    content_size: int
    sheets: list[CatalogSheetPreview]
    requires_sheet_choice: bool
    max_size_bytes: int


def normalize_identifier(value: Any) -> str:
    """Normalize an OE/MPN without ever coercing it to a number."""

    raw = _cell_text(value)
    primary = _IDENTIFIER_SPLIT_RE.split(raw, maxsplit=1)[0]
    return normalize_oem_identifier(primary)


def preview_catalog_xlsx(
    content: bytes,
    *,
    filename: str,
    sample_size: int = 5,
) -> CatalogWorkbookPreview:
    """Inspect workbook structure without choosing a sheet or writing data.

    Cost-labelled cells are intentionally excluded from preview rows while the
    cost privacy mode is undecided. The column name may be shown so an operator
    understands why it cannot be mapped, but its values never cross the API.
    """

    _validate_xlsx_container(content)
    try:
        workbook = load_workbook(
            BytesIO(content),
            read_only=True,
            data_only=True,
            keep_links=False,
        )
    except (InvalidFileException, OSError, ValueError, KeyError) as exc:
        raise CatalogImportError("Файл не является валидным XLSX") from exc

    previews: list[CatalogSheetPreview] = []
    try:
        for worksheet in workbook.worksheets:
            iterator = worksheet.iter_rows(values_only=True)
            try:
                header_values = next(iterator)
            except StopIteration:
                previews.append(
                    CatalogSheetPreview(
                        name=str(worksheet.title),
                        row_count=0,
                        headers=[],
                        suggested_mapping={},
                        mapping_error="В листе нет строк",
                        sample_rows=[],
                        is_catalog_candidate=False,
                    )
                )
                continue
            if len(header_values) > MAX_COLUMNS:
                previews.append(
                    CatalogSheetPreview(
                        name=str(worksheet.title),
                        row_count=0,
                        headers=[],
                        suggested_mapping={},
                        mapping_error=(
                            f"Слишком много колонок: {len(header_values)} > "
                            f"{MAX_COLUMNS}"
                        ),
                        sample_rows=[],
                        is_catalog_candidate=False,
                    )
                )
                continue
            headers = _unique_headers(header_values)
            mapping_error: str | None = None
            suggested_mapping: dict[str, str] = {}
            try:
                resolved = resolve_column_mapping(headers)
                suggested_mapping = {
                    field: headers[index] for field, index in resolved.items()
                }
            except CatalogImportError as exc:
                mapping_error = str(exc)

            row_count = 0
            sample_rows: list[dict[str, Any]] = []
            for values in iterator:
                if not any(value not in (None, "") for value in values):
                    continue
                row_count += 1
                if len(sample_rows) >= sample_size:
                    continue
                sample_rows.append(
                    {
                        header: _json_safe(
                            values[index] if index < len(values) else None
                        )
                        for index, header in enumerate(headers)
                        if not is_raw_cost_label(header)
                    }
                )
            previews.append(
                CatalogSheetPreview(
                    name=str(worksheet.title),
                    row_count=row_count,
                    headers=headers,
                    suggested_mapping=suggested_mapping,
                    mapping_error=mapping_error,
                    sample_rows=sample_rows,
                    is_catalog_candidate=mapping_error is None and row_count > 0,
                )
            )
    finally:
        workbook.close()

    if not previews:
        raise CatalogImportError("В XLSX нет листов")
    return CatalogWorkbookPreview(
        filename=filename[:255] or "catalog.xlsx",
        content_sha256=hashlib.sha256(content).hexdigest(),
        content_size=len(content),
        sheets=previews,
        requires_sheet_choice=True,
        max_size_bytes=MAX_XLSX_BYTES,
    )


def _validate_xlsx_container(content: bytes) -> None:
    if not content:
        raise CatalogImportError("Файл пуст")
    if len(content) > MAX_XLSX_BYTES:
        raise CatalogImportError(f"Файл больше {MAX_XLSX_BYTES // 1024 // 1024} MB")
    try:
        with ZipFile(BytesIO(content)) as archive:
            uncompressed_size = sum(entry.file_size for entry in archive.infolist())
    except (BadZipFile, OSError) as exc:
        raise CatalogImportError("Файл не является валидным XLSX") from exc
    if uncompressed_size > MAX_UNCOMPRESSED_XLSX_BYTES:
        raise CatalogImportError("Распакованный XLSX слишком велик")


def parse_catalog_xlsx(
    content: bytes,
    *,
    explicit_mapping: dict[str, str] | None = None,
    sheet_name: str | None = None,
    allow_encrypted_cost_input: bool = False,
    characteristics_config: CharacteristicsConfig | None = None,
) -> ParsedCatalog:
    _validate_xlsx_container(content)
    try:
        workbook = load_workbook(
            BytesIO(content),
            read_only=True,
            data_only=True,
            keep_links=False,
        )
    except (InvalidFileException, OSError, ValueError, KeyError) as exc:
        raise CatalogImportError("Файл не является валидным XLSX") from exc

    try:
        if sheet_name is not None:
            if sheet_name not in workbook.sheetnames:
                raise CatalogImportError(f"Лист {sheet_name!r} не найден")
            sheet = workbook[sheet_name]
        else:
            sheet = _default_sheet(workbook)
        sheet_title = str(sheet.title)
        iterator = sheet.iter_rows(values_only=True)
        try:
            header_values = next(iterator)
        except StopIteration as exc:
            raise CatalogImportError("В XLSX нет строк") from exc
        if len(header_values) > MAX_COLUMNS:
            raise CatalogImportError(
                f"Слишком много колонок: {len(header_values)} > {MAX_COLUMNS}"
            )
        headers = _unique_headers(header_values)
        resolved = resolve_column_mapping(headers, explicit_mapping)
        oe_column_asserts_identity = _oe_column_asserts_identity(
            headers,
            resolved,
            explicit_mapping=explicit_mapping,
        )
        config = characteristics_config or load_characteristics_config(
            backend_config_path(get_settings().catalog_characteristics_path)
        )
        columns = characteristic_column_pairs(headers, config)
        recognized_names: dict[str, int] = {}
        unrecognized_names: dict[str, int] = {}
        anomaly_counts: dict[str, int] = {}
        self_reference_total = 0
        cost_indexes = [
            index for index, header in enumerate(headers) if is_raw_cost_label(header)
        ]
        if len(cost_indexes) > 1:
            raise CatalogImportError(
                "Найдено несколько колонок себестоимости; оставьте одну"
            )
        cost_index = cost_indexes[0] if cost_indexes else None

        rows: list[ParsedCatalogRow] = []
        issues: list[ImportIssue] = []
        sensitive_costs: dict[int, Decimal] = {}
        seen_skus: set[str] = set()
        seen_oe_rows: dict[str, ParsedCatalogRow] = {}
        blocked_oe_collisions: set[str] = set()
        identity_collision_rows = 0
        total_rows = 0
        for source_row, values in enumerate(iterator, start=2):
            if source_row - 1 > MAX_ROWS:
                raise CatalogImportError(f"В XLSX больше {MAX_ROWS} строк")
            if not any(value not in (None, "") for value in values):
                continue
            total_rows += 1
            raw_row = {
                header: _json_safe(values[index] if index < len(values) else None)
                for index, header in enumerate(headers)
                if not is_raw_cost_label(header)
            }
            # Characteristic names are counted for every non-empty row, even one
            # that later fails validation: the report describes the workbook, not
            # the subset that survived.
            collected = collect_characteristics(values, columns.pairs)
            for name in collected:
                target = (
                    recognized_names
                    if normalize_characteristic_name(name) in config.rules_by_name
                    else unrecognized_names
                )
                target[name] = target.get(name, 0) + 1
            try:
                raw_cost = (
                    values[cost_index]
                    if cost_index is not None and cost_index < len(values)
                    else None
                )
                parsed_cost = _optional_unit_cost(raw_cost)
                if parsed_cost is not None and not allow_encrypted_cost_input:
                    raise SensitiveCatalogImportBlocked(
                        "Себестоимость заполнена, но защищённый серверный импорт не включён"
                    )
                parsed, extraction = _parse_row(
                    source_row,
                    values,
                    resolved,
                    raw_row,
                    collected=collected,
                    config=config,
                    columns_balanced=columns.balanced,
                    oe_column_asserts_identity=oe_column_asserts_identity,
                )
                for code in extraction.anomalies:
                    anomaly_counts[code] = anomaly_counts.get(code, 0) + 1
                self_reference_total += extraction.dropped_self_references
                if parsed.sku in seen_skus:
                    raise CatalogImportError(f"Дублирующийся SKU: {parsed.sku}")
                seen_skus.add(parsed.sku)
                if (
                    parsed.identity_status == "OE_CONFIRMED"
                    and parsed.oe_norm in blocked_oe_collisions
                ):
                    # Preserve the source row for operator review.  A repeated
                    # normalized OE is not safe identity evidence, but dropping
                    # the row hides a real customer SKU and makes the import
                    # less complete.  The UNRESOLVED state is fail-closed at
                    # every matching/pricing boundary.
                    parsed = replace(
                        parsed,
                        identity_status="UNRESOLVED",
                        identity_reason="NORMALIZED_OE_COLLISION",
                    )
                    rows.append(parsed)
                    identity_collision_rows += 1
                    if parsed_cost is not None:
                        sensitive_costs[source_row] = parsed_cost
                    continue
                previous = (
                    seen_oe_rows.get(parsed.oe_norm)
                    if parsed.identity_status == "OE_CONFIRMED"
                    else None
                )
                if previous is not None:
                    previous_index = rows.index(previous)
                    rows[previous_index] = replace(
                        previous,
                        identity_status="UNRESOLVED",
                        identity_reason="NORMALIZED_OE_COLLISION",
                    )
                    seen_oe_rows.pop(parsed.oe_norm, None)
                    blocked_oe_collisions.add(parsed.oe_norm)
                    parsed = replace(
                        parsed,
                        identity_status="UNRESOLVED",
                        identity_reason="NORMALIZED_OE_COLLISION",
                    )
                    rows.append(parsed)
                    identity_collision_rows += 2
                    if parsed_cost is not None:
                        sensitive_costs[source_row] = parsed_cost
                    continue
                if parsed.identity_status == "OE_CONFIRMED":
                    seen_oe_rows[parsed.oe_norm] = parsed
                rows.append(parsed)
                if parsed.internal_code_raw and not parsed.internal_code_norm:
                    # Строку не роняем: она годная, просто по внутреннему коду
                    # не свяжется. Молчать нельзя — это выглядело бы как связь.
                    issues.append(
                        ImportIssue(
                            source_row,
                            "INTERNAL_CODE_IGNORED",
                            "внутренний код не распознан как наш: "
                            f"{parsed.internal_code_raw}",
                        )
                    )
                if parsed_cost is not None:
                    sensitive_costs[source_row] = parsed_cost
            except SensitiveCatalogImportBlocked:
                raise
            except CatalogImportError as exc:
                issues.append(ImportIssue(source_row, "INVALID_ROW", str(exc)))
    finally:
        workbook.close()

    if total_rows == 0:
        raise CatalogImportError("В XLSX нет товарных строк")
    accepted_source_rows = {row.source_row for row in rows}
    column_mapping = {field: headers[index] for field, index in resolved.items()}
    if cost_index is not None:
        column_mapping["cost"] = headers[cost_index]
    return ParsedCatalog(
        rows=rows,
        issues=issues,
        column_mapping=column_mapping,
        total_rows=total_rows,
        sensitive_costs={
            row: value
            for row, value in sensitive_costs.items()
            if row in accepted_source_rows
        },
        characteristics_report={
            "schema_version": config.schema_version,
            "method_version": config.method_version,
            "config_sha256": config.source_sha256,
            "sheet": sheet_title,
            "self_references_dropped": self_reference_total,
            "column_pairs": len(columns.pairs),
            "columns_balanced": columns.balanced,
            "recognized": dict(sorted(recognized_names.items())),
            "unrecognized": dict(sorted(unrecognized_names.items())),
            "anomalies": dict(sorted(anomaly_counts.items())),
            "identity_collision_rows": identity_collision_rows,
            "rows_with_part_numbers": sum(1 for row in rows if row.part_numbers_norm),
            "part_numbers_total": sum(len(row.part_numbers_norm) for row in rows),
            "rows_with_applicability_brand": sum(
                1 for row in rows if row.applicability_brands
            ),
            "rows_with_applicability_model": sum(
                1 for row in rows if row.applicability_models
            ),
        },
    )


def _default_sheet(workbook: Any) -> Any:
    """Pick the sheet to import, refusing to guess between equal candidates.

    A real Prom export was observed with three sheets where the *active* one was
    a seventeen-row scratch sheet carrying the same headers as the real
    forty-nine-hundred-row product sheet.  Trusting ``workbook.active`` there
    imports 16 products out of 4901 and reports success, which is the worst
    possible outcome: silent, plausible, and wrong.  When more than one sheet
    could be a catalog, the caller must say which.
    """

    candidates: list[str] = []
    for worksheet in workbook.worksheets:
        try:
            header_values = next(worksheet.iter_rows(values_only=True))
        except StopIteration:
            continue
        if len(header_values) > MAX_COLUMNS:
            continue
        try:
            resolve_column_mapping(_unique_headers(header_values))
        except CatalogImportError:
            continue
        candidates.append(worksheet.title)
    if len(candidates) > 1:
        raise CatalogImportError(
            "Каталог найден на нескольких листах ("
            + ", ".join(repr(title) for title in candidates)
            + "); укажите лист явно"
        )
    if candidates:
        return workbook[candidates[0]]
    return workbook.active


def resolve_column_mapping(
    headers: list[str], explicit_mapping: dict[str, str] | None = None
) -> dict[str, int]:
    normalized_headers: dict[str, list[int]] = {}
    for index, header in enumerate(headers):
        normalized_headers.setdefault(_normalize_header(header), []).append(index)

    explicit = explicit_mapping or {}
    unknown_fields = set(explicit) - set(FIELD_ALIASES)
    if unknown_fields:
        raise CatalogImportError(
            "Неизвестные поля mapping: " + ", ".join(sorted(unknown_fields))
        )
    if "cost" in explicit:
        raise CatalogImportError(
            "Raw cost import is disabled until cost privacy mode is approved"
        )

    result = _prom_export_mapping(headers) if not explicit else {}
    for field, aliases in FIELD_ALIASES.items():
        if field in result:
            continue
        candidates = (explicit[field],) if field in explicit else aliases
        matched: list[int] = []
        for candidate in candidates:
            matched.extend(normalized_headers.get(_normalize_header(candidate), []))
        matched = sorted(set(matched))
        if len(matched) > 1:
            raise CatalogImportError(
                f"Колонка {field!r} определилась неоднозначно; передайте mapping явно"
            )
        if matched:
            result[field] = matched[0]
        elif field in explicit:
            raise CatalogImportError(
                f"Колонка {explicit[field]!r} для поля {field!r} не найдена"
            )

    missing = REQUIRED_FIELDS - set(result)
    if missing:
        raise CatalogImportError(
            "Не найдены обязательные колонки: " + ", ".join(sorted(missing))
        )
    return result


async def import_catalog_xlsx(
    session: AsyncSession,
    *,
    workspace_id: UUID,
    filename: str,
    content: bytes,
    explicit_mapping: dict[str, str] | None = None,
    sheet_name: str | None = None,
    store_id: UUID | None = None,
    user_id: UUID,
    settings: Settings | None = None,
) -> CatalogImportBatch:
    """Parse and atomically persist one immutable catalog snapshot."""
    selected = settings or get_settings()
    allow_encrypted_cost_input = False
    if (
        CostPrivacyMode(selected.cost_privacy_mode)
        == CostPrivacyMode.SERVER_SIDE_ENCRYPTED
    ):
        try:
            selected.cost_keyring
        except ValueError:
            pass
        else:
            allow_encrypted_cost_input = True
    parsed = await asyncio.to_thread(
        parse_catalog_xlsx,
        content,
        explicit_mapping=explicit_mapping,
        sheet_name=sheet_name,
        allow_encrypted_cost_input=allow_encrypted_cost_input,
    )
    request_fingerprint = _catalog_import_fingerprint(
        content_sha256=hashlib.sha256(content).hexdigest(),
        sheet_name=str(parsed.characteristics_report.get("sheet") or ""),
        column_mapping=parsed.column_mapping,
    )
    existing = await _existing_idempotent_import(
        session,
        workspace_id=workspace_id,
        request_fingerprint=request_fingerprint,
    )
    if existing is not None:
        return existing
    if parsed.sensitive_costs:
        try:
            require_server_cost_input_allowed(
                next(iter(parsed.sensitive_costs.values())), settings=selected
            )
        except RuntimeError as exc:
            raise CatalogImportError(
                "Защищённый импорт себестоимости недоступен"
            ) from exc
    now = datetime.now(UTC)
    row_outcomes = catalog_row_outcomes(parsed)
    rejected_row_count = sum(
        row["terminal_status"] == "REJECTED_NOT_IMPORTABLE" for row in row_outcomes
    )
    batch = CatalogImportBatch(
        workspace_id=workspace_id,
        filename=filename[:255] or "catalog.xlsx",
        content_sha256=hashlib.sha256(content).hexdigest(),
        request_fingerprint=request_fingerprint,
        content_size=len(content),
        status="running",
        column_mapping=parsed.column_mapping,
        total_rows=parsed.total_rows,
        imported_rows=len(parsed.rows),
        rejected_rows=rejected_row_count,
        error_log=[asdict(issue) for issue in parsed.issues[:MAX_ERROR_LOG]],
        row_outcomes_contract_version=CATALOG_ROW_OUTCOMES_CONTRACT_VERSION,
        row_outcomes_sha256=_canonical_json_sha256(row_outcomes),
        row_outcomes=row_outcomes,
        characteristics_report=parsed.characteristics_report,
        started_at=now,
    )
    session.add(batch)
    await session.flush()
    items = [
        CatalogItem(
            workspace_id=workspace_id,
            import_batch_id=batch.id,
            store_id=store_id,
            **asdict(row),
        )
        for row in parsed.rows
    ]
    session.add_all(items)
    await session.flush()
    await upsert_xlsx_products(session, batch=batch, items=items)
    for item, row in zip(items, parsed.rows, strict=True):
        cost = parsed.sensitive_costs.get(row.source_row)
        if cost is not None:
            add_encrypted_cost_record(
                session,
                workspace_id=workspace_id,
                catalog_item_id=item.id,
                user_id=user_id,
                cost=cost,
                reason=f"encrypted XLSX import row {row.source_row}",
                settings=selected,
            )
    batch.status = (
        "failed"
        if not parsed.rows
        else "partial"
        if rejected_row_count
        else "completed"
    )
    batch.finished_at = datetime.now(UTC)
    await session.commit()
    await session.refresh(batch)
    return batch


def _catalog_import_fingerprint(
    *,
    content_sha256: str,
    sheet_name: str,
    column_mapping: Mapping[str, str],
) -> str:
    payload = json.dumps(
        {
            "content_sha256": content_sha256,
            "sheet_name": sheet_name,
            "column_mapping": dict(sorted(column_mapping.items())),
        },
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    )
    return hashlib.sha256(payload.encode()).hexdigest()


async def _existing_idempotent_import(
    session: AsyncSession,
    *,
    workspace_id: UUID,
    request_fingerprint: str,
) -> CatalogImportBatch | None:
    """Serialize identical PostgreSQL requests and reuse a terminal snapshot."""

    scalar = getattr(session, "scalar", None)
    execute = getattr(session, "execute", None)
    if scalar is None or execute is None:
        # Lightweight unit-test sessions exercise encryption/persistence only.
        return None
    bind = getattr(session, "bind", None)
    dialect = getattr(getattr(bind, "dialect", None), "name", None)
    if dialect == "postgresql":
        await execute(
            text(
                "SELECT pg_advisory_xact_lock("
                "hashtextextended(:fingerprint, 0))"
            ),
            {"fingerprint": request_fingerprint},
        )
    existing = await scalar(
        select(CatalogImportBatch)
        .where(
            CatalogImportBatch.workspace_id == workspace_id,
            CatalogImportBatch.request_fingerprint == request_fingerprint,
            CatalogImportBatch.status.in_(("completed", "partial")),
        )
        .order_by(CatalogImportBatch.created_at.desc(), CatalogImportBatch.id.desc())
        .limit(1)
    )
    if existing is not None:
        await session.commit()
    return existing


async def get_import_batch(
    session: AsyncSession, *, workspace_id: UUID, batch_id: UUID
) -> CatalogImportBatch | None:
    return await session.scalar(
        select(CatalogImportBatch).where(
            CatalogImportBatch.id == batch_id,
            CatalogImportBatch.workspace_id == workspace_id,
        )
    )


async def build_catalog_terminal_manifest(
    session: AsyncSession,
    *,
    workspace_id: UUID,
    batch_id: UUID,
    run_id: UUID | None = None,
) -> dict[str, Any] | None:
    """Build a tenant-scoped source-row manifest, optionally through replay."""

    batch = await get_import_batch(
        session,
        workspace_id=workspace_id,
        batch_id=batch_id,
    )
    if batch is None:
        return None
    items = list(
        (
            await session.scalars(
                select(CatalogItem)
                .where(
                    CatalogItem.workspace_id == workspace_id,
                    CatalogItem.import_batch_id == batch.id,
                )
                .order_by(CatalogItem.source_row, CatalogItem.id)
            )
        ).all()
    )
    item_by_source_row = {item.source_row: item for item in items}
    pricing_run: PricingRun | None = None
    run_item_by_catalog_item: dict[UUID, PricingRunItem] = {}
    recommendation_by_run_item: dict[UUID, PricingRecommendation] = {}
    if run_id is not None:
        pricing_run = await session.scalar(
            select(PricingRun).where(
                PricingRun.id == run_id,
                PricingRun.workspace_id == workspace_id,
                PricingRun.import_batch_id == batch.id,
            )
        )
        if pricing_run is None:
            return None
        run_items = list(
            (
                await session.scalars(
                    select(PricingRunItem)
                    .where(PricingRunItem.pricing_run_id == pricing_run.id)
                    .order_by(PricingRunItem.membership_position, PricingRunItem.id)
                )
            ).all()
        )
        run_item_by_catalog_item = {
            run_item.catalog_item_id: run_item for run_item in run_items
        }
        if run_items:
            recommendations = list(
                (
                    await session.scalars(
                        select(PricingRecommendation).where(
                            PricingRecommendation.pricing_run_item_id.in_(
                                [run_item.id for run_item in run_items]
                            )
                        )
                    )
                ).all()
            )
            recommendation_by_run_item = {
                recommendation.pricing_run_item_id: recommendation
                for recommendation in recommendations
            }

    contract_version = batch.row_outcomes_contract_version
    stored_hash = batch.row_outcomes_sha256
    if contract_version == CATALOG_ROW_OUTCOMES_CONTRACT_VERSION:
        rows = [dict(row) for row in (batch.row_outcomes or [])]
        recomputed_hash = _canonical_json_sha256(rows)
        source = "PINNED_ROW_OUTCOMES"
        hash_verified = recomputed_hash == stored_hash
    else:
        accepted_rows = {
            item.source_row: {
                "source_row": item.source_row,
                "terminal_status": "IMPORTED",
                "sku": item.sku,
                "oe_norm": item.oe_norm,
                "reason_codes": [],
                "details": [],
            }
            for item in items
        }
        rejected_rows: dict[int, dict[str, Any]] = {}
        for raw in batch.error_log or []:
            if not isinstance(raw, Mapping):
                continue
            try:
                source_row = int(raw.get("row"))
            except (TypeError, ValueError):
                continue
            row = rejected_rows.setdefault(
                source_row,
                {
                    "source_row": source_row,
                    "terminal_status": "REJECTED_NOT_IMPORTABLE",
                    "sku": None,
                    "oe_norm": None,
                    "reason_codes": [],
                    "details": [],
                },
            )
            code = str(raw.get("code") or "").strip()
            message = str(raw.get("message") or "").strip()
            if code and code not in row["reason_codes"]:
                row["reason_codes"].append(code)
            if message:
                row["details"].append(message[:1000])
        overlap = set(accepted_rows) & set(rejected_rows)
        known = {
            **{key: value for key, value in accepted_rows.items() if key not in overlap},
            **{key: value for key, value in rejected_rows.items() if key not in overlap},
        }
        rows = [
            {"source_ordinal": ordinal, **known[source_row]}
            for ordinal, source_row in enumerate(sorted(known), start=1)
        ]
        recomputed_hash = _canonical_json_sha256(rows)
        stored_hash = None
        source = "LEGACY_RECONSTRUCTED"
        hash_verified = False

    terminal_statuses = {"IMPORTED", "REJECTED_NOT_IMPORTABLE", "FAILED", "CANCELLED"}
    ordinals = [row.get("source_ordinal") for row in rows]
    import_complete = (
        len(rows) == batch.total_rows
        and ordinals == list(range(1, batch.total_rows + 1))
        and len({row.get("source_row") for row in rows}) == len(rows)
        and all(row.get("terminal_status") in terminal_statuses for row in rows)
        and all(
            row.get("terminal_status") != "REJECTED_NOT_IMPORTABLE"
            or bool(row.get("reason_codes"))
            for row in rows
        )
        and (
            hash_verified
            if contract_version == CATALOG_ROW_OUTCOMES_CONTRACT_VERSION
            else len(items) + batch.rejected_rows == batch.total_rows
            and len(batch.error_log or []) == batch.rejected_rows
        )
    )
    rendered_rows: list[dict[str, Any]] = []
    for row in rows:
        rendered = dict(row)
        item = item_by_source_row.get(int(row.get("source_row") or 0))
        rendered["catalog_item_id"] = str(item.id) if item is not None else None
        if pricing_run is not None:
            run_item = (
                run_item_by_catalog_item.get(item.id) if item is not None else None
            )
            recommendation = (
                recommendation_by_run_item.get(run_item.id)
                if run_item is not None
                else None
            )
            not_imported = row.get("terminal_status") != "IMPORTED"
            has_terminal_recommendation = bool(
                run_item is not None
                and run_item.status in {"calculated", "manual_review"}
                and recommendation is not None
            )
            has_terminal_failure = bool(
                run_item is not None
                and run_item.status in {"failed", "cancelled"}
            )
            identity_blocked_recommendation = bool(
                recommendation is not None
                and not recommendation_price_identity_allowed(
                    item, recommendation.action
                )
            )
            rendered.update(
                {
                    "pricing_run_item_id": (
                        str(run_item.id) if run_item is not None else None
                    ),
                    "pricing_terminal_status": (
                        run_item.status.upper()
                        if run_item is not None
                        else "NOT_APPLICABLE_NOT_IMPORTED"
                        if not_imported
                        else "MISSING_RUN_ITEM"
                    ),
                    "pricing_terminal": (
                        has_terminal_recommendation
                        or has_terminal_failure
                        or (run_item is None and not_imported)
                    ),
                    "pricing_terminal_reason": (
                        IDENTITY_BLOCKED_RECOMMENDATION_ACTION
                        if identity_blocked_recommendation
                        else "RECOMMENDATION_PERSISTED"
                        if has_terminal_recommendation
                        else "RUN_ITEM_TERMINAL_FAILURE"
                        if has_terminal_failure
                        else "NOT_APPLICABLE_NOT_IMPORTED"
                        if run_item is None and not_imported
                        else "RECOMMENDATION_MISSING"
                        if run_item is not None
                        and run_item.status in {"calculated", "manual_review"}
                        else "RUN_ITEM_NON_TERMINAL"
                        if run_item is not None
                        else "RUN_ITEM_MISSING"
                    ),
                    "pricing_error": run_item.error if run_item is not None else None,
                    "recommendation_id": (
                        str(recommendation.id) if recommendation is not None else None
                    ),
                    "recommendation_action": (
                        (
                            "MANUAL_REVIEW"
                            if identity_blocked_recommendation
                            else recommendation.action
                        )
                        if recommendation is not None
                        else None
                    ),
                }
            )
        rendered_rows.append(rendered)
    replay_nonterminal_count = (
        sum(not bool(row.get("pricing_terminal")) for row in rendered_rows)
        if pricing_run is not None
        else None
    )
    replay_missing_item_count = (
        sum(
            row.get("terminal_status") == "IMPORTED"
            and row.get("pricing_run_item_id") is None
            for row in rendered_rows
        )
        if pricing_run is not None
        else None
    )
    pricing_replay_complete = (
        pricing_run is not None
        and pricing_run.status in {"completed", "partial", "failed", "cancelled"}
        and replay_nonterminal_count == 0
        and replay_missing_item_count == 0
        and len(run_item_by_catalog_item) == len(items)
    )
    complete = import_complete and (
        pricing_replay_complete if pricing_run is not None else True
    )
    return {
        "manifest_version": (
            "catalog-pricing-replay-terminal-manifest-v2"
            if pricing_run is not None
            else "catalog-terminal-manifest-v1"
        ),
        "batch_id": batch.id,
        "pricing_run_id": pricing_run.id if pricing_run is not None else None,
        "pricing_run_status": pricing_run.status if pricing_run is not None else None,
        "catalog_content_sha256": batch.content_sha256,
        "row_outcomes_contract_version": contract_version,
        "source": source,
        "expected_rows": batch.total_rows,
        "manifest_rows": len(rows),
        "import_manifest_complete": import_complete,
        "pricing_replay_complete": (
            pricing_replay_complete if pricing_run is not None else None
        ),
        "pricing_terminal_rows": (
            sum(bool(row.get("pricing_terminal")) for row in rendered_rows)
            if pricing_run is not None
            else None
        ),
        "pricing_nonterminal_rows": replay_nonterminal_count,
        "pricing_missing_run_items": replay_missing_item_count,
        "complete": complete,
        "verification_status": "VERIFIED" if complete else "NOT_PROVEN",
        "row_outcomes_sha256": stored_hash,
        "recomputed_sha256": recomputed_hash,
        "hash_verified": hash_verified,
        "silent_loss_count": max(0, batch.total_rows - len(rows))
        + (replay_missing_item_count or 0),
        "replay_manifest_sha256": (
            _canonical_json_sha256(rendered_rows)
            if pricing_run is not None
            else None
        ),
        "rows": rendered_rows,
    }


async def list_import_batches(
    session: AsyncSession,
    *,
    workspace_id: UUID,
    limit: int,
    offset: int,
) -> tuple[list[CatalogImportBatch], int]:
    where = CatalogImportBatch.workspace_id == workspace_id
    total = int(
        await session.scalar(select(func.count(CatalogImportBatch.id)).where(where))
        or 0
    )
    batches = list(
        (
            await session.scalars(
                select(CatalogImportBatch)
                .where(where)
                .order_by(
                    CatalogImportBatch.created_at.desc(), CatalogImportBatch.id.desc()
                )
                .limit(limit)
                .offset(offset)
            )
        ).all()
    )
    return batches, total


async def list_catalog_items(
    session: AsyncSession,
    *,
    workspace_id: UUID,
    batch_id: UUID | None,
    limit: int,
    offset: int,
) -> tuple[list[CatalogItem], int]:
    conditions = [CatalogItem.workspace_id == workspace_id]
    if batch_id is not None:
        conditions.append(CatalogItem.import_batch_id == batch_id)
    total = int(
        await session.scalar(select(func.count(CatalogItem.id)).where(*conditions)) or 0
    )
    items = list(
        (
            await session.scalars(
                select(CatalogItem)
                .where(*conditions)
                .order_by(CatalogItem.source_row, CatalogItem.id)
                .limit(limit)
                .offset(offset)
            )
        ).all()
    )
    return items, total


def parse_mapping_json(value: str | None) -> dict[str, str] | None:
    if value is None or not value.strip():
        return None
    try:
        parsed = json.loads(value)
    except json.JSONDecodeError as exc:
        raise CatalogImportError("mapping должен быть JSON-объектом") from exc
    if not isinstance(parsed, dict) or not all(
        isinstance(key, str) and isinstance(item, str) for key, item in parsed.items()
    ):
        raise CatalogImportError("mapping должен содержать пары строка-строка")
    return parsed


def _parse_row(
    source_row: int,
    values: tuple[Any, ...],
    mapping: dict[str, int],
    raw_row: dict[str, Any],
    *,
    collected: Mapping[str, Sequence[str]],
    config: CharacteristicsConfig,
    columns_balanced: bool,
    oe_column_asserts_identity: bool,
) -> tuple[ParsedCatalogRow, CharacteristicsExtraction]:
    def get(field: str) -> Any:
        index = mapping.get(field)
        return values[index] if index is not None and index < len(values) else None

    oe_raw = _cell_text(get("oe"))
    oe_norm = normalize_identifier(oe_raw)
    if oe_raw and oe_column_asserts_identity and len(oe_norm) < 3:
        raise CatalogImportError("невалидный OE/OEM")
    if len(oe_norm) > 255:
        raise CatalogImportError("OE/OEM длиннее 255 символов")
    name = _required_text(get("name"), "название")
    category = _required_text(get("category"), "категория")[:255]
    price = _positive_decimal(get("price"), "цена")
    mpn_raw = _cell_text(get("mpn"))
    mpn_norm = normalize_identifier(mpn_raw)
    internal_code_raw = _cell_text(get("internal_code"))
    internal_code_norm = (
        normalize_identifier(internal_code_raw)
        if is_internal_catalog_code(internal_code_raw)
        else ""
    )
    sku = _cell_text(get("sku")) or (
        f"OE-{oe_norm}-{source_row}"
        if oe_norm
        else f"MPN-{mpn_norm}-{source_row}"
        if mpn_norm
        else f"ROW-{source_row}"
    )
    if len(sku) > 255:
        raise CatalogImportError("SKU длиннее 255 символов")
    currency = _normalize_currency(get("currency"))
    stock_qty = _optional_nonnegative_decimal(get("stock_qty"), "остаток")
    stock_age = _optional_nonnegative_decimal(get("stock_age_days"), "возраст запаса")
    expected = _optional_nonnegative_decimal(get("expected_units_sold"), "продажи")
    units_sold_30d = _optional_nonnegative_decimal(
        get("units_sold_30d"), "продажи за 30 дней"
    )
    units_sold_60d = _optional_nonnegative_decimal(
        get("units_sold_60d"), "продажи за 60 дней"
    )
    units_sold_90d = _optional_nonnegative_decimal(
        get("units_sold_90d"), "продажи за 90 дней"
    )
    days_since_last_sale = _optional_nonnegative_decimal(
        get("days_since_last_sale"), "дней с последней продажи"
    )
    historical_monthly_units = _optional_nonnegative_decimal(
        get("historical_monthly_units"), "средние продажи в месяц"
    )
    views_30d = _optional_nonnegative_decimal(get("views_30d"), "просмотры за 30 дней")
    conversion_rate_proxy = _optional_ratio(get("conversion_rate_proxy"), "конверсия")
    priority = _optional_positive_decimal(
        get("manual_priority"), "приоритет"
    ) or Decimal("1")
    product_url = _optional_text(get("product_url"))
    if product_url and not product_url.casefold().startswith(("https://", "http://")):
        raise CatalogImportError(
            "ссылка на товар должна начинаться с http:// или https://"
        )
    # The row's own codes are excluded from its cross list: a number cannot be
    # its own cross, and ``cross_links`` rejects such a pair by check constraint.
    extraction = extract_characteristics(
        collected,
        config,
        self_numbers=(oe_raw, sku, mpn_raw),
        columns_balanced=columns_balanced,
    )
    if oe_norm and oe_column_asserts_identity:
        identity_status = "OE_CONFIRMED"
        identity_reason = "CUSTOMER_OE_COLUMN"
    elif mpn_norm:
        identity_status = "MPN_ONLY"
        identity_reason = "CUSTOMER_MPN_COLUMN"
    elif extraction.part_numbers_norm:
        identity_status = "MPN_ONLY"
        identity_reason = "CUSTOMER_PART_NUMBER_LIST"
    else:
        identity_status = "UNRESOLVED"
        identity_reason = "CUSTOMER_IDENTITY_MISSING"
    return ParsedCatalogRow(
        source_row=source_row,
        sku=sku,
        oe_raw=oe_raw,
        oe_norm=oe_norm,
        mpn_raw=mpn_raw,
        mpn_norm=mpn_norm,
        internal_code_raw=internal_code_raw,
        internal_code_norm=internal_code_norm,
        name=name,
        category=category,
        brand=_optional_text(get("brand"), max_length=255),
        description=_optional_text(get("description")),
        product_url=product_url,
        current_price=price,
        currency=currency,
        is_available=_optional_bool(get("available")),
        stock_status=_stock_status(get("stock_status")),
        stock_qty=stock_qty,
        stock_age_days=stock_age,
        expected_units_sold=expected,
        units_sold_30d=units_sold_30d,
        units_sold_60d=units_sold_60d,
        units_sold_90d=units_sold_90d,
        days_since_last_sale=days_since_last_sale,
        historical_monthly_units=historical_monthly_units,
        views_30d=views_30d,
        conversion_rate_proxy=conversion_rate_proxy,
        cost=None,
        manual_priority=priority,
        raw_row=raw_row,
        part_numbers_raw=list(extraction.part_numbers_raw),
        part_numbers_norm=list(extraction.part_numbers_norm),
        applicability_brands=list(extraction.applicability_brands),
        applicability_models=list(extraction.applicability_models),
        characteristics_raw=dict(extraction.characteristics_raw),
        identity_status=identity_status,
        identity_reason=identity_reason,
    ), extraction


def _oe_column_asserts_identity(
    headers: Sequence[str],
    mapping: Mapping[str, int],
    *,
    explicit_mapping: Mapping[str, str] | None,
) -> bool:
    """Whether the customer/operator explicitly labelled the mapped value as OE.

    Prom's canonical ``Код_товару`` column is deliberately *not* such a claim:
    the supplied workbook mixes internal KEMP shelf codes and real part numbers
    in it.  Treating every value as an OE both invents identity and turns
    repeated internal codes into false normalized-OE collisions.
    """

    oe_index = mapping.get("oe")
    if oe_index is None:
        return False
    if explicit_mapping is not None and "oe" in explicit_mapping:
        return True
    normalized = _normalize_header(headers[oe_index])
    return bool(
        re.search(r"(?:^| )(?:oe|oem|ое)(?: |$)", normalized)
        or "оригинальн" in normalized
        or "оригінальн" in normalized
    )


def _prom_export_mapping(headers: list[str]) -> dict[str, int]:
    """Recognize Yuri's canonical Prom.ua export without an explicit JSON map."""
    exact = {_normalize_header(header): index for index, header in enumerate(headers)}
    required = {
        "sku": "унікальний ідентифікатор",
        "oe": "код товару",
        "name": "назва позиції",
        "category": "назва групи",
        "price": "ціна",
    }
    if not all(header in exact for header in required.values()):
        return {}
    optional = {
        "mpn": "номер пристрою mpn",
        "currency": "валюта",
        "available": "наявність",
        "brand": "виробник",
        "description": "опис",
        "product_url": "продукт на сайті",
        "stock_qty": "кількість",
    }
    result = {field: exact[header] for field, header in required.items()}
    result.update(
        {field: exact[header] for field, header in optional.items() if header in exact}
    )
    return result


def _unique_headers(values: tuple[Any, ...]) -> list[str]:
    result: list[str] = []
    seen: dict[str, int] = {}
    for index, value in enumerate(values, start=1):
        base = _cell_text(value) or f"column_{index}"
        count = seen.get(base, 0) + 1
        seen[base] = count
        result.append(base if count == 1 else f"{base}#{count}")
    return result


def _normalize_header(value: str) -> str:
    return _HEADER_CLEAN_RE.sub(" ", value.casefold().replace("_", " ")).strip()


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return format(value, ".15g")
    return str(value).strip()


def _required_text(value: Any, label: str) -> str:
    normalized = _cell_text(value)
    if not normalized:
        raise CatalogImportError(f"пустое поле: {label}")
    return normalized


def _optional_text(value: Any, *, max_length: int | None = None) -> str | None:
    normalized = _cell_text(value)
    if not normalized:
        return None
    return normalized[:max_length] if max_length else normalized


def _decimal(value: Any, label: str) -> Decimal:
    if isinstance(value, bool) or value is None:
        raise CatalogImportError(f"невалидное число: {label}")
    if isinstance(value, (int, float, Decimal)):
        raw = str(value)
    else:
        raw = _cell_text(value).replace("\u00a0", "").replace(" ", "")
        raw = re.sub(r"(?i)(uah|грн|гривен|гривня|₴)", "", raw)
        if "," in raw and "." in raw:
            if raw.rfind(",") > raw.rfind("."):
                raw = raw.replace(".", "").replace(",", ".")
            else:
                raw = raw.replace(",", "")
        elif "," in raw:
            raw = raw.replace(",", ".")
    try:
        result = Decimal(raw)
    except (InvalidOperation, ValueError) as exc:
        raise CatalogImportError(f"невалидное число: {label}") from exc
    if not result.is_finite():
        raise CatalogImportError(f"невалидное число: {label}")
    return result


def _positive_decimal(value: Any, label: str) -> Decimal:
    result = _decimal(value, label)
    if result <= 0:
        raise CatalogImportError(f"{label} должна быть > 0")
    return result.quantize(Decimal("0.01"))


def _optional_positive_decimal(value: Any, label: str) -> Decimal | None:
    if value is None or not _cell_text(value):
        return None
    return _positive_decimal(value, label)


def _optional_nonnegative_decimal(value: Any, label: str) -> Decimal | None:
    if value is None or not _cell_text(value):
        return None
    result = _decimal(value, label)
    if result < 0:
        raise CatalogImportError(f"{label} не может быть отрицательным")
    return result


def _optional_ratio(value: Any, label: str) -> Decimal | None:
    result = _optional_nonnegative_decimal(value, label)
    if result is not None and result > 1:
        raise CatalogImportError(f"{label} должна быть от 0 до 1")
    return result


def _optional_unit_cost(value: Any) -> Decimal | None:
    if value is None or not _cell_text(value):
        return None
    result = _decimal(value, "себестоимость")
    if result <= 0:
        raise CatalogImportError("себестоимость должна быть > 0")
    if result.as_tuple().exponent < -2:
        raise CatalogImportError("себестоимость поддерживает не более 2 знаков")
    if result >= Decimal("1000000000000"):
        raise CatalogImportError("себестоимость выше допустимого диапазона")
    return result.quantize(Decimal("0.01"))


def _normalize_currency(value: Any) -> str:
    raw = _cell_text(value).casefold()
    if not raw or raw in {"uah", "грн", "₴", "гривня", "гривень"}:
        return "UAH"
    normalized = raw.upper()
    if len(normalized) != 3 or not normalized.isalpha():
        raise CatalogImportError("валюта должна быть ISO-4217 кодом")
    return normalized


def _optional_bool(value: Any) -> bool | None:
    if value is None or not _cell_text(value):
        return None
    if isinstance(value, bool):
        return value
    raw = _cell_text(value).casefold()
    if raw in {
        "1",
        "true",
        "yes",
        "y",
        "да",
        "так",
        "есть",
        "в наличии",
        "в наявності",
        "available",
        "+",
        "!",
    }:
        return True
    if raw in {
        "0",
        "false",
        "no",
        "n",
        "нет",
        "ні",
        "нет в наличии",
        "немає",
        "unavailable",
        "-",
    }:
        return False
    raise CatalogImportError("неизвестное значение наличия")


def _stock_status(value: Any) -> str:
    raw = _cell_text(value).casefold().replace("-", "_").replace(" ", "_")
    if not raw:
        return "unknown"
    aliases = {
        "fresh": "fresh",
        "свежий": "fresh",
        "свіжий": "fresh",
        "ходовой": "fresh",
        "stale": "stale",
        "лежалый": "stale",
        "залежалый": "stale",
        "залежалий": "stale",
        "dead_stock": "dead_stock",
        "deadstock": "dead_stock",
        "неликвид": "dead_stock",
        "неліквід": "dead_stock",
        "unknown": "unknown",
        "неизвестно": "unknown",
    }
    try:
        return aliases[raw]
    except KeyError as exc:
        raise CatalogImportError(
            "статус должен быть fresh/stale/dead_stock/unknown"
        ) from exc


def _json_safe(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


__all__ = [
    "CatalogImportError",
    "CatalogSheetPreview",
    "CatalogWorkbookPreview",
    "ImportIssue",
    "ParsedCatalog",
    "ParsedCatalogRow",
    "SensitiveCatalogImportBlocked",
    "get_import_batch",
    "import_catalog_xlsx",
    "list_catalog_items",
    "list_import_batches",
    "normalize_identifier",
    "parse_catalog_xlsx",
    "parse_mapping_json",
    "preview_catalog_xlsx",
    "resolve_column_mapping",
]
