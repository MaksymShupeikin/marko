"""Filter-faithful CSV/XLSX export for pricing recommendations.

The export contract deliberately contains no cost or inventory-cost fields.
Those remain unavailable until the workspace cost-privacy mode is approved.
"""

from __future__ import annotations

import csv
from dataclasses import dataclass
from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO, StringIO
from typing import Any, Literal
from uuid import UUID

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from marko.infrastructure.db.models import CatalogItem, MarketObservation
from marko.services.catalog_identity_safety import is_internal_catalog_code
from marko.services.market_collection import _validated_listing_url
from marko.services.pricing_runs import (
    IDENTITY_BLOCKED_RECOMMENDATION_ACTION,
    customer_identity_query_from_fields,
    list_recommendations,
    recommendation_price_identity_allowed,
)
from metis.pricing import normalize_oe


MAX_RECOMMENDATION_EXPORT_ROWS = 5_000
MAX_SOURCE_URLS_PER_ROW = 10


class RecommendationExportError(ValueError):
    def __init__(self, code: str, message: str) -> None:
        self.code = code
        super().__init__(message)


@dataclass(frozen=True, slots=True)
class RecommendationExport:
    content: bytes
    media_type: str
    filename: str
    row_count: int
    run_id: UUID | None


async def export_recommendations(
    session: AsyncSession,
    *,
    workspace_id: UUID,
    export_format: Literal["csv", "xlsx"],
    run_id: UUID | None,
    action: str | None,
    confidence_grade: str | None,
    category: str | None,
    queue: str,
    priority_score_type: str | None,
    confidence_min: Decimal | None,
    confidence_max: Decimal | None,
    sort: str,
) -> RecommendationExport:
    rows, total, resolved_run_id, _action_counts = await list_recommendations(
        session,
        workspace_id=workspace_id,
        run_id=run_id,
        action=action,
        confidence_grade=confidence_grade,
        category=category,
        queue=queue,
        priority_score_type=priority_score_type,
        confidence_min=confidence_min,
        confidence_max=confidence_max,
        sort=sort,
        limit=MAX_RECOMMENDATION_EXPORT_ROWS,
        offset=0,
    )
    if total > MAX_RECOMMENDATION_EXPORT_ROWS:
        raise RecommendationExportError(
            "RECOMMENDATION_EXPORT_LIMIT_EXCEEDED",
            "Активные фильтры возвращают "
            f"{total} строк; сузьте выборку до "
            f"{MAX_RECOMMENDATION_EXPORT_ROWS}.",
        )
    source_urls = await _source_urls_by_recommendation(session, rows)
    payload = [
        _export_row(
            recommendation,
            item,
            source_urls=source_urls.get(recommendation.id, ()),
        )
        for recommendation, item in rows
    ]
    suffix = str(resolved_run_id or "empty")
    if export_format == "csv":
        content = _to_csv(payload)
        return RecommendationExport(
            content=content,
            media_type="text/csv; charset=utf-8",
            filename=f"marko-recommendations-{suffix}.csv",
            row_count=len(payload),
            run_id=resolved_run_id,
        )
    content = _to_xlsx(payload)
    return RecommendationExport(
        content=content,
        media_type=(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ),
        filename=f"marko-recommendations-{suffix}.xlsx",
        row_count=len(payload),
        run_id=resolved_run_id,
    )


async def _source_urls_by_recommendation(
    session: AsyncSession,
    rows: list[tuple[Any, CatalogItem]],
) -> dict[UUID, tuple[str, ...]]:
    ids: set[UUID] = set()
    recommendation_ids: dict[UUID, list[UUID]] = {}
    for recommendation, _item in rows:
        ordered: list[UUID] = []
        for raw in (
            list(recommendation.evidence_observation_ids or [])
            + list(recommendation.kemp_reference_observation_ids or [])
        ):
            try:
                observation_id = UUID(str(raw))
            except (TypeError, ValueError):
                continue
            if observation_id not in ordered:
                ordered.append(observation_id)
                ids.add(observation_id)
        recommendation_ids[recommendation.id] = ordered
    if not ids:
        return {recommendation_id: () for recommendation_id in recommendation_ids}
    observations = {
        observation.id: observation
        for observation in (
            await session.scalars(
                select(MarketObservation).where(MarketObservation.id.in_(ids))
            )
        ).all()
    }
    result: dict[UUID, tuple[str, ...]] = {}
    for recommendation_id, ordered_ids in recommendation_ids.items():
        urls: list[str] = []
        for observation_id in ordered_ids:
            observation = observations.get(observation_id)
            if observation is None:
                continue
            safe_url, _reason = _validated_listing_url(observation.url)
            if safe_url and safe_url not in urls:
                urls.append(safe_url)
            if len(urls) == MAX_SOURCE_URLS_PER_ROW:
                break
        result[recommendation_id] = tuple(urls)
    return result


def _export_row(
    recommendation: Any,
    item: CatalogItem,
    *,
    source_urls: tuple[str, ...],
) -> dict[str, Any]:
    identity = _export_identity_fields(item)
    identity_blocked = not recommendation_price_identity_allowed(
        item, recommendation.action
    )
    export_action = "MANUAL_REVIEW" if identity_blocked else recommendation.action
    export_reason_codes = list(recommendation.reason_codes or [])
    if (
        identity_blocked
        and IDENTITY_BLOCKED_RECOMMENDATION_ACTION not in export_reason_codes
    ):
        export_reason_codes.append(IDENTITY_BLOCKED_RECOMMENDATION_ACTION)
    money = lambda value: _format_money(  # noqa: E731
        value,
        currency=recommendation.currency,
        price_tick=recommendation.price_tick,
        price_tick_version=recommendation.price_tick_version,
    )
    trace = getattr(recommendation, "calculation_trace", {})
    advisory = (
        trace.get("advisory_decision")
        if isinstance(trace, dict)
        else None
    )
    if not isinstance(advisory, dict):
        advisory = {}
    if identity_blocked:
        # A stale MPN_ONLY recommendation may carry an old customer advisory
        # price in its trace.  It is still auditable in the database, but an
        # export is an operational handoff and must be safe by construction.
        advisory = {}
    return {
        "sku": item.sku,
        # ``oe`` is reserved for an asserted vehicle/OEM number.  A legacy
        # MPN_ONLY row may still carry a private KEMP shelf code in
        # ``CatalogItem.oe_norm``; exporting that value under the OE header
        # would make a correct search look semantically wrong to the
        # operator.  The public manufacturer number and retrieval key are
        # exported separately below.
        "oe": identity["oe"],
        "mpn": identity["mpn"],
        "search_identity": identity["search_identity"],
        "identity_status": identity["identity_status"],
        "name": item.name,
        "category": item.category,
        "action": export_action,
        "current_price": money(recommendation.current_price),
        "fair_price": money(None if identity_blocked else recommendation.fair_price),
        "recommended_price": money(
            None if identity_blocked else recommendation.recommended_price
        ),
        "absolute_recommended_change": money(
            None
            if identity_blocked
            else recommendation.absolute_recommended_change
        ),
        "percentage_recommended_change": (
            None
            if identity_blocked or recommendation.percentage_recommended_change is None
            else str(recommendation.percentage_recommended_change)
        ),
        "customer_advisory_action": advisory.get("action", ""),
        "customer_advisory_price": money(
            _optional_decimal(advisory.get("recommended_price"))
        ),
        "customer_target_band_low": money(
            _optional_decimal(advisory.get("target_band_low"))
        ),
        "customer_target_band_high": money(
            _optional_decimal(advisory.get("target_band_high"))
        ),
        "automatic_price_application": (
            "false"
            if identity_blocked
            or advisory.get("automatic_price_application") is False
            else ""
        ),
        "confidence": str(recommendation.confidence),
        "confidence_grade": recommendation.confidence_grade,
        "reason_codes": "; ".join(export_reason_codes),
        "source_urls": list(source_urls),
        "computed_at": recommendation.computed_at.isoformat(),
    }


def _export_identity_fields(item: CatalogItem) -> dict[str, str]:
    """Expose identity without relabelling a private code as OE.

    The pricing run already freezes this namespace decision.  Reusing the
    same pure resolver here keeps CSV/XLSX output from becoming a second,
    weaker identity implementation.
    """

    raw_status = getattr(item, "identity_status", None)
    status = str(raw_status or "UNRESOLVED").strip().upper()
    # Small in-process test/replay adapters predating ``identity_status`` are
    # treated as explicit OE only when they omit the field.  Real CatalogItem
    # rows always carry the persisted status and therefore remain fail-closed
    # when unresolved.
    if raw_status is None:
        status = "OE_CONFIRMED" if getattr(item, "oe_norm", None) else "UNRESOLVED"

    def public(value: Any) -> str:
        normalized = normalize_oe(str(value or ""))
        if not normalized or is_internal_catalog_code(normalized):
            return ""
        return normalized

    oe = public(getattr(item, "oe_norm", None))
    mpn = public(getattr(item, "mpn_norm", None))
    part_numbers = tuple(
        value
        for raw in (getattr(item, "part_numbers_norm", None) or ())
        if (value := public(raw))
    )
    search_identity = customer_identity_query_from_fields(
        identity_status=status,
        oe_norm=oe,
        mpn_norm=mpn,
        part_numbers_norm=part_numbers,
    )
    return {
        "oe": oe if status == "OE_CONFIRMED" else "",
        "mpn": mpn,
        "search_identity": search_identity,
        "identity_status": status,
    }


def _optional_decimal(value: Any) -> Decimal | None:
    if value is None or str(value).strip() == "":
        return None
    try:
        parsed = Decimal(str(value))
    except ArithmeticError:
        return None
    return parsed if parsed.is_finite() else None


def _format_money(
    value: Decimal | None,
    *,
    currency: str,
    price_tick: Decimal,
    price_tick_version: str | None = None,
) -> str:
    if value is None:
        return ""
    tick = Decimal(price_tick)
    if tick <= 0:
        tick = Decimal("0.01")
    rounded = (Decimal(value) / tick).quantize(
        Decimal("1"),
        rounding=ROUND_HALF_UP,
    ) * tick
    # NUMERIC(14,4) pads the configured integer tick `1` to `1.0000`.
    # The named tick contract, not database padding, controls presentation.
    digits = (
        0
        if price_tick_version == "uah-integer-v1"
        else max(0, -tick.as_tuple().exponent)
    )
    return f"{rounded:.{digits}f} {currency.strip().upper()}"


def _to_csv(rows: list[dict[str, Any]]) -> bytes:
    output = StringIO(newline="")
    headers = _headers(include_source_columns=False)
    writer = csv.DictWriter(output, fieldnames=headers, extrasaction="ignore")
    writer.writeheader()
    for row in rows:
        flattened = dict(row)
        flattened["source_urls"] = "; ".join(row["source_urls"])
        writer.writerow(flattened)
    return ("\ufeff" + output.getvalue()).encode("utf-8")


def _to_xlsx(rows: list[dict[str, Any]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Recommendations"
    headers = _headers(include_source_columns=True)
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        values = [
            row.get(header, "")
            if not header.startswith("source_url_")
            else _source_at(row["source_urls"], header)
            for header in headers
        ]
        sheet.append(values)
        row_number = sheet.max_row
        for column, header in enumerate(headers, 1):
            if not header.startswith("source_url_"):
                continue
            cell = sheet.cell(row=row_number, column=column)
            if cell.value:
                cell.hyperlink = str(cell.value)
                cell.style = "Hyperlink"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def _headers(*, include_source_columns: bool) -> list[str]:
    headers = [
        "sku",
        "oe",
        "mpn",
        "search_identity",
        "identity_status",
        "name",
        "category",
        "action",
        "current_price",
        "fair_price",
        "recommended_price",
        "absolute_recommended_change",
        "percentage_recommended_change",
        "customer_advisory_action",
        "customer_advisory_price",
        "customer_target_band_low",
        "customer_target_band_high",
        "automatic_price_application",
        "confidence",
        "confidence_grade",
        "reason_codes",
    ]
    if include_source_columns:
        headers.extend(
            f"source_url_{index}"
            for index in range(1, MAX_SOURCE_URLS_PER_ROW + 1)
        )
    else:
        headers.append("source_urls")
    headers.append("computed_at")
    return headers


def _source_at(urls: list[str], header: str) -> str:
    index = int(header.rsplit("_", 1)[1]) - 1
    return urls[index] if index < len(urls) else ""


__all__ = [
    "MAX_RECOMMENDATION_EXPORT_ROWS",
    "MAX_SOURCE_URLS_PER_ROW",
    "RecommendationExport",
    "RecommendationExportError",
    "export_recommendations",
]
